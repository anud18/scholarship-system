"""統計教育部博士生獎學金-各系所獲獎生人數 — per-department 教育部 counts.

One row per 系所 of the office's reference form, grouped under its 教育部統計處
學科標準分類 field (``MOE_DEPARTMENT_FIELD_LAYOUT``):

- 申請人數: students who applied for a 教育部 sub-type (any sub-type whose
  funding agency is 教育部 — ``moe_1w`` / ``moe_2w``).
- 通過甄選人數(正取): students the 分發名單 allocated a 教育部 sub-type.
- 加碼補助人數: left blank for the office to fill in by hand.

Departments missing from the form (a new 系所, or one with no 教育部 awardee
the year the form was drawn) are appended under ``UNMAPPED_FIELD_LABEL`` so no
student is silently dropped. Cell counts are distinct students.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.scholarship import ScholarshipType
from app.models.student import Department
from app.services.export_table_chassis import apply_xlsx_borders
from app.services.moe_department_stats_layout import MOE_DEPARTMENT_FIELD_LAYOUT
from app.services.stats_applicant_source import load_applicant_rows
from app.utils.excel_safety import sanitize_excel_cell
from app.utils.export_download import sanitise_filename_part
from app.utils.sub_type_agency import sub_type_agency_label

MOE_AGENCY_LABEL = "教育部"
REPORT_NAME = "各系所獲獎生人數"
UNMAPPED_FIELD_LABEL = "未對應領域"
UNKNOWN_DEPARTMENT_LABEL = "未知系所"
TOTAL_ROW_LABEL = "合計"
HEADERS = (
    "領域(請填寫教育部統計處大專校院學科標準分類11領域)",
    "系所完整名稱",
    "申請人數",
    "通過甄選人數(正取)",
    "加碼補助人數",
)

# Reference form styling (統計教育部博士生獎學金-各系所獲獎生人數_1151008.xlsx).
_FONT_NAME = "標楷體"
_FONT_SIZE = 14
_FIELD_HEADER_FONT_SIZE = 11
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
_TITLE_ROW_HEIGHT = 30.75
_HEADER_ROW_HEIGHT = 58.5
_BODY_ROW_HEIGHT = 19.5
_COLUMN_WIDTHS = (24.5, 47.875, 13, 13, 13)

_WHITESPACE = re.compile(r"\s+")


@dataclass(frozen=True)
class DepartmentRecord:
    """One counted student: who, which 系所 (resolved name or bare code), which sub-types."""

    user_id: int
    department_name: Optional[str]
    department_code: Optional[str]
    sub_type_codes: Tuple[str, ...]


@dataclass(frozen=True)
class DepartmentStatsRow:
    field: str
    department: str
    applicants: int
    awardees: int


@dataclass(frozen=True)
class MoeDepartmentStatsExport:
    payload: bytes
    filename: str
    applicant_count: int
    awardee_count: int


def is_moe_sub_type(sub_type_code: Optional[str]) -> bool:
    return sub_type_agency_label(sub_type_code or "") == MOE_AGENCY_LABEL


def _normalise_name(name: str) -> str:
    """Match key for a 系所 name: no whitespace, half-width parentheses."""
    return _WHITESPACE.sub("", name).replace("（", "(").replace("）", ")")


_ROW_BY_NAME: Dict[str, Tuple[str, str]] = {
    _normalise_name(alias): (field, label)
    for field, rows in MOE_DEPARTMENT_FIELD_LAYOUT
    for label, aliases in rows
    for alias in aliases
}


def department_code_from_data(student_data: Optional[Mapping[str, Any]]) -> Optional[str]:
    """``std_depno`` first — the same std-first precedence the 獲獎人數 sheets use."""
    data = student_data or {}
    code = data.get("std_depno") or data.get("trm_depno")
    return str(code).strip() if code else None


def department_record(
    user_id: int,
    student_data: Optional[Mapping[str, Any]],
    sub_type_codes: Iterable[Optional[str]],
    name_by_code: Mapping[str, str],
) -> DepartmentRecord:
    """Resolve the student's 系所 name from ``departments`` (by code).

    Falls back to the snapshot's ``trm_depname`` only when the code came from
    the same term record, so a name never pairs with another 系所's code.
    """
    data = student_data or {}
    code = department_code_from_data(data)
    name = name_by_code.get(code) if code else None
    if name is None and code and str(data.get("trm_depno") or "").strip() == code:
        name = (data.get("trm_depname") or "").strip() or None
    return DepartmentRecord(
        user_id=user_id,
        department_name=name,
        department_code=code,
        sub_type_codes=tuple(c for c in sub_type_codes if c),
    )


def _row_key(record: DepartmentRecord) -> Tuple[str, str]:
    if record.department_name:
        mapped = _ROW_BY_NAME.get(_normalise_name(record.department_name))
        if mapped:
            return mapped
        return (UNMAPPED_FIELD_LABEL, record.department_name)
    if record.department_code:
        return (UNMAPPED_FIELD_LABEL, f"{UNKNOWN_DEPARTMENT_LABEL}（{record.department_code}）")
    return (UNMAPPED_FIELD_LABEL, UNKNOWN_DEPARTMENT_LABEL)


def _moe_users_by_row(records: Sequence[DepartmentRecord]) -> Dict[Tuple[str, str], Set[int]]:
    users: Dict[Tuple[str, str], Set[int]] = {}
    for record in records:
        if any(is_moe_sub_type(code) for code in record.sub_type_codes):
            users.setdefault(_row_key(record), set()).add(record.user_id)
    return users


def tally_department_counts(
    applicants: Sequence[DepartmentRecord], awardees: Sequence[DepartmentRecord]
) -> List[DepartmentStatsRow]:
    """Every layout row (zeros included — it is a fixed form), then unmapped 系所 by name.

    Only records carrying a 教育部 sub-type count. No 合計 row here — the
    workbook adds it.
    """
    applicant_users = _moe_users_by_row(applicants)
    awardee_users = _moe_users_by_row(awardees)

    keys = [(field, label) for field, rows in MOE_DEPARTMENT_FIELD_LAYOUT for label, _ in rows]
    layout_keys = set(keys)
    keys += sorted({*applicant_users, *awardee_users} - layout_keys, key=lambda k: k[1])
    return [
        DepartmentStatsRow(
            field=field,
            department=label,
            applicants=len(applicant_users.get((field, label), ())),
            awardees=len(awardee_users.get((field, label), ())),
        )
        for field, label in keys
    ]


def _field_spans(rows: Sequence[DepartmentStatsRow]) -> List[Tuple[str, int, int]]:
    """(field, first index, last index) for each run of consecutive rows sharing a field."""
    spans: List[Tuple[str, int, int]] = []
    for index, row in enumerate(rows):
        if spans and spans[-1][0] == row.field:
            spans[-1] = (row.field, spans[-1][1], index)
        else:
            spans.append((row.field, index, index))
    return spans


def build_workbook(*, title: str, rows: Sequence[DepartmentStatsRow]) -> bytes:
    """Render the reference form: merged 領域 cells, one row per 系所, then 合計."""
    wb = Workbook()
    ws = wb.active
    ws.title = REPORT_NAME
    column_count = len(HEADERS)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    justified = Alignment(horizontal="justify", vertical="center", wrap_text=True)
    body_font = Font(name=_FONT_NAME, size=_FONT_SIZE)

    ws.cell(row=1, column=1, value=sanitize_excel_cell(title))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    ws.cell(row=1, column=1).font = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True)
    ws.cell(row=1, column=1).alignment = centered

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(name=_FONT_NAME, size=_FIELD_HEADER_FONT_SIZE if col == 1 else _FONT_SIZE)
        cell.alignment = centered
        if col >= 3:
            cell.fill = _HEADER_FILL

    first_body_row = 3
    for offset, row in enumerate(rows):
        r = first_body_row + offset
        # Column E (加碼補助人數) stays blank — the office fills it in by hand.
        for col, value in enumerate((row.field, row.department, row.applicants, row.awardees), start=1):
            cell = ws.cell(row=r, column=col, value=sanitize_excel_cell(value))
            cell.font = body_font
            cell.alignment = justified if col <= 2 else centered
    for _field, start, end in _field_spans(rows):
        if end > start:
            ws.merge_cells(start_row=first_body_row + start, start_column=1, end_row=first_body_row + end, end_column=1)

    total_row = first_body_row + len(rows)
    total_font = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True)
    ws.cell(row=total_row, column=1, value=TOTAL_ROW_LABEL)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    ws.cell(row=total_row, column=3, value=sum(row.applicants for row in rows))
    ws.cell(row=total_row, column=4, value=sum(row.awardees for row in rows))
    for col in range(1, column_count + 1):
        ws.cell(row=total_row, column=col).font = total_font
        ws.cell(row=total_row, column=col).alignment = centered

    ws.row_dimensions[1].height = _TITLE_ROW_HEIGHT
    ws.row_dimensions[2].height = _HEADER_ROW_HEIGHT
    for r in range(first_body_row, total_row + 1):
        ws.row_dimensions[r].height = _BODY_ROW_HEIGHT
    for col, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = width
    apply_xlsx_borders(ws, max_row=total_row, max_col=column_count, first_row=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def load_scholarship(db: AsyncSession, scholarship_type_id: int) -> ScholarshipType:
    """The scholarship with its sub-type configs, to gate on 教育部 sub-types."""
    # scalar_one(): callers reach here through finalized rankings that FK onto
    # this scholarship type, so a miss is a broken invariant, not a 404.
    return (
        await db.execute(
            select(ScholarshipType)
            .options(selectinload(ScholarshipType.sub_type_configs))
            .where(ScholarshipType.id == scholarship_type_id)
        )
    ).scalar_one()


def scholarship_has_moe_sub_type(scholarship: ScholarshipType) -> bool:
    return any(is_moe_sub_type(config.sub_type_code) for config in scholarship.sub_type_configs)


async def _department_names(
    db: AsyncSession, student_data_rows: Iterable[Optional[Mapping[str, Any]]]
) -> Dict[str, str]:
    codes = {code for code in (department_code_from_data(data) for data in student_data_rows) if code}
    if not codes:
        return {}
    rows = (await db.execute(select(Department.code, Department.name).where(Department.code.in_(codes)))).all()
    return {code: name for code, name in rows}


async def build_moe_department_award_stats_export(
    db: AsyncSession,
    *,
    scholarship: ScholarshipType,
    academic_year: int,
    semester: str,
    awards: Sequence[Tuple[int, Optional[Mapping[str, Any]], Optional[str]]],
    title_prefix: str,
) -> MoeDepartmentStatsExport:
    """Render the sheet for one scholarship / year / semester.

    ``awards`` is one ``(user_id, student_data, allocated_sub_type)`` per
    allocated 分發名單 slot; applicants are loaded here with the same filter as
    統計各學院學生申請人數.
    """
    applicant_rows = await load_applicant_rows(
        db, scholarship_type_id=scholarship.id, academic_year=academic_year, semester=semester
    )
    name_by_code = await _department_names(
        db, [row.student_data for row in applicant_rows] + [data for _, data, _ in awards]
    )
    applicants = [department_record(r.user_id, r.student_data, r.sub_type_codes, name_by_code) for r in applicant_rows]
    awardees = [department_record(uid, data, (sub_type,), name_by_code) for uid, data, sub_type in awards]
    rows = tally_department_counts(applicants, awardees)

    scholarship_name = scholarship.name or scholarship.code
    report_title = f"統計{MOE_AGENCY_LABEL}{scholarship_name}-{REPORT_NAME}"
    return MoeDepartmentStatsExport(
        payload=build_workbook(title=f"{title_prefix}{report_title}", rows=rows),
        filename=f"{sanitise_filename_part(f'{title_prefix}{report_title}')}.xlsx",
        applicant_count=sum(row.applicants for row in rows),
        awardee_count=sum(row.awardees for row in rows),
    )
