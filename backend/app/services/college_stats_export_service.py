"""Per-college student count Excel exports — 統計各學院學生申請人數 / 獲獎人數.

Pure logic — the endpoints load and authorize the source rows (applications for
the 申請人數 sheet; allocated distribution / roster items for the 獲獎人數
sheet), then hand this module plain ``StatsRecord`` values. Both sheets share
the office's reference layout: 編號 / 學院 / one column per funding agency
(國科會, 教育部, …), one row per college in the office's fixed order.

Columns are derived from the scholarship's own sub-type configuration, so a
scholarship without the 國科會/教育部 sub-types still gets correct columns.
Sub-types sharing a funding agency (``moe_1w`` / ``moe_2w`` → 教育部) collapse
into one column; any other sub-type gets its own column under its configured
name. A student is counted once per column even if they picked (or were
allocated) several sub-types that map to it.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Dict, FrozenSet, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.core.college_mappings import COLLEGE_MAPPINGS
from app.utils.academic_period import format_academic_period
from app.services.export_table_chassis import apply_xlsx_borders
from app.utils.excel_safety import sanitize_excel_cell
from app.utils.sub_type_agency import sub_type_agency_label

APPLICANT_REPORT_NAME = "統計各學院學生申請人數"
AWARD_REPORT_NAME = "統計各學院學生獲獎人數"

# (row label, SIS academy codes counted in that row) in the office's reference
# order. Labels are the colleges' full official names (the academies table only
# carries SIS short names). Code 8 is the 陽明-campus code for the same 人社院
# as A, so both feed one row.
COLLEGE_STATS_LAYOUT: Tuple[Tuple[str, FrozenSet[str]], ...] = (
    ("人文藝術與社會學院", frozenset({"A", "8"})),
    ("工學院", frozenset({"I"})),
    ("工程生物科學學院", frozenset({"B"})),
    ("客家文化學院", frozenset({"K"})),
    ("理學院", frozenset({"S"})),
    ("資訊學院", frozenset({"C"})),
    ("電機學院", frozenset({"E"})),
    ("管理學院", frozenset({"M"})),
    ("光電學院", frozenset({"O"})),
    ("科技法律學院", frozenset({"L"})),
    ("國際半導體產業學院", frozenset({"D"})),
    ("智慧科學暨綠能學院", frozenset({"G"})),
    ("牙醫學院", frozenset({"2"})),
    ("生命科學院", frozenset({"7"})),
    ("生物醫學暨工程學院", frozenset({"6"})),
    ("醫學院", frozenset({"1"})),
    ("藥物科學院", frozenset({"5"})),
    ("護理學院", frozenset({"3"})),
    ("產學創新研究學院", frozenset({"F"})),
)

UNKNOWN_COLLEGE_LABEL = "未知學院"
TOTAL_ROW_LABEL = "合計"
APPLICANT_COUNT_LABEL = "申請人數"
AWARD_COUNT_LABEL = "獲獎人數"
# Key of the lone count column a scholarship without sub-types gets; its label
# is the caller's (申請人數 / 獲獎人數).
SINGLE_COLUMN_KEY = "__all__"
UNSPECIFIED_SUB_TYPE_COLUMN = "未指定類別"

# Placeholder sub-type codes used by scholarships that have no real sub-types.
_NON_SUB_TYPE_CODES = frozenset({"", "general", "default"})

_FONT_NAME = "微軟正黑體"
_FONT_SIZE = 16
_ROW_HEIGHT = 21
_INDEX_COL_WIDTH = 8
_COLLEGE_COL_WIDTH = 30
_COUNT_COL_WIDTH = 16


@dataclass(frozen=True)
class StatsRecord:
    """The three facts the tally needs from one counted row (application / award)."""

    user_id: int
    academy_code: Optional[str]
    sub_type_codes: Tuple[str, ...]


@dataclass(frozen=True)
class StatsColumn:
    key: str
    label: str


@dataclass(frozen=True)
class StatsRow:
    label: str
    counts: Tuple[int, ...]


def stats_period_label(academic_year: int, semester: Optional[str]) -> str:
    """Title period: ``115學年度`` for a whole year (semester None), else ``115學年度第一學期``."""
    if semester is None:
        return f"{academic_year}學年度"
    return format_academic_period(academic_year, semester)


def _normalize_code(code: Optional[str]) -> str:
    return (code or "").lower().strip()


def _column_key(sub_type_code: str) -> str:
    return sub_type_agency_label(sub_type_code) or sub_type_code


def _real_sub_types(codes: Iterable[str]) -> List[str]:
    return [c for c in (_normalize_code(x) for x in codes) if c not in _NON_SUB_TYPE_CODES]


def build_stats_columns(
    sub_type_configs: Sequence[Tuple[str, str]],
    records: Sequence[StatsRecord],
    known_names: Optional[Mapping[str, str]] = None,
    single_column_label: str = APPLICANT_COUNT_LABEL,
) -> List[StatsColumn]:
    """Columns for the count matrix.

    ``sub_type_configs`` is the scholarship's active ``(sub_type_code, name)``
    list in display order. Codes that appear on applications but are not
    active (e.g. a deactivated sub-type) are appended so no applicant is
    silently dropped, labelled from ``known_names`` (every configured
    sub-type, active or not). A scholarship with no real sub-types gets a
    single ``single_column_label`` column.
    """
    names = {_normalize_code(code): name for code, name in (known_names or {}).items()}
    names.update({_normalize_code(code): name for code, name in sub_type_configs})
    ordered_codes = _real_sub_types(code for code, _ in sub_type_configs)
    observed = sorted({c for r in records for c in _real_sub_types(r.sub_type_codes)} - set(ordered_codes))

    columns: List[StatsColumn] = []
    seen: Set[str] = set()
    for code in ordered_codes + observed:
        key = _column_key(code)
        if key in seen:
            continue
        seen.add(key)
        columns.append(StatsColumn(key=key, label=sub_type_agency_label(code) or names.get(code) or code))

    if not columns:
        return [StatsColumn(key=SINGLE_COLUMN_KEY, label=single_column_label)]
    if any(not _real_sub_types(r.sub_type_codes) for r in records):
        columns.append(StatsColumn(key=UNSPECIFIED_SUB_TYPE_COLUMN, label=UNSPECIFIED_SUB_TYPE_COLUMN))
    return columns


def _record_column_keys(record: StatsRecord, column_keys: Set[str]) -> Set[str]:
    if SINGLE_COLUMN_KEY in column_keys:
        return {SINGLE_COLUMN_KEY}
    keys = {_column_key(c) for c in _real_sub_types(record.sub_type_codes)}
    return keys or {UNSPECIFIED_SUB_TYPE_COLUMN}


def _row_label_for_extra_code(code: str) -> str:
    if not code:
        return UNKNOWN_COLLEGE_LABEL
    return COLLEGE_MAPPINGS.get(code) or f"{UNKNOWN_COLLEGE_LABEL}（{code}）"


def tally_college_counts(records: Sequence[StatsRecord], columns: Sequence[StatsColumn]) -> List[StatsRow]:
    """One row per layout college, then any colleges outside the layout, then 合計.

    Every layout row is emitted even at zero (the sheet is a fixed form). Cell
    counts are distinct students, so a student with two applications in one
    cell is counted once; 合計 is the column sum.
    """
    column_keys = [c.key for c in columns]
    key_set = set(column_keys)
    row_by_code = {code: label for label, codes in COLLEGE_STATS_LAYOUT for code in codes}

    users: Dict[Tuple[str, str], Set[int]] = {}
    extra_labels: Dict[str, str] = {}
    for record in records:
        code = (record.academy_code or "").strip()
        row_label = row_by_code.get(code)
        if row_label is None:
            row_label = _row_label_for_extra_code(code)
            extra_labels[code] = row_label
        for key in _record_column_keys(record, key_set):
            users.setdefault((row_label, key), set()).add(record.user_id)

    row_labels = [label for label, _ in COLLEGE_STATS_LAYOUT]
    row_labels += [extra_labels[code] for code in sorted(extra_labels, key=lambda c: (not c, c))]

    rows = [StatsRow(label, tuple(len(users.get((label, key), ())) for key in column_keys)) for label in row_labels]
    # Column sums, so the printed form always adds up (a student whose two
    # applications carry different colleges appears in both rows).
    totals = tuple(sum(row.counts[i] for row in rows) for i in range(len(column_keys)))
    return rows + [StatsRow(TOTAL_ROW_LABEL, totals)]


def build_workbook(
    *, title: str, columns: Sequence[StatsColumn], rows: Sequence[StatsRow], sheet_name: str = APPLICANT_REPORT_NAME
) -> bytes:
    """Render the count matrix in the reference sheet's style."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    font = Font(name=_FONT_NAME, size=_FONT_SIZE)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    total_cols = 2 + len(columns)

    ws.cell(row=1, column=1, value=sanitize_excel_cell(title))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws.cell(row=1, column=1).font = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True)
    ws.cell(row=1, column=1).alignment = centered

    headers = ["編號", "學院"] + [c.label for c in columns]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=sanitize_excel_cell(header))
        cell.font = font
        cell.alignment = left if col == 2 else centered

    for offset, row in enumerate(rows):
        r = 3 + offset
        is_total = row.label == TOTAL_ROW_LABEL
        index_value = None if is_total else offset + 1
        values = [index_value, row.label, *row.counts]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=sanitize_excel_cell(value))
            cell.font = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=is_total)
            cell.alignment = left if col == 2 else centered

    last_row = 2 + len(rows)
    for r in range(1, last_row + 1):
        ws.row_dimensions[r].height = _ROW_HEIGHT
    ws.column_dimensions["A"].width = _INDEX_COL_WIDTH
    ws.column_dimensions["B"].width = _COLLEGE_COL_WIDTH
    for col in range(3, total_cols + 1):
        ws.column_dimensions[ws.cell(row=2, column=col).column_letter].width = _COUNT_COL_WIDTH
    apply_xlsx_borders(ws, max_row=last_row, max_col=total_cols, first_row=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_award_stats_workbook(
    *,
    title: str,
    records: Sequence[StatsRecord],
    sub_type_configs: Sequence[Tuple[str, str]],
    known_names: Mapping[str, str],
) -> bytes:
    """統計各學院學生獲獎人數 — one record per awarded (allocated) student slot."""
    columns = build_stats_columns(
        sub_type_configs, records, known_names=known_names, single_column_label=AWARD_COUNT_LABEL
    )
    return build_workbook(
        title=title,
        columns=columns,
        rows=tally_college_counts(records, columns),
        sheet_name=AWARD_REPORT_NAME,
    )
