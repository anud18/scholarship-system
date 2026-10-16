"""Endpoint tests for the admin 統計各學院學生申請人數 export.

Counts are asserted by parsing the workbook, not by scanning response bytes.
"""

import io
from urllib.parse import unquote

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import get_current_user
from app.main import app
from app.models.application import Application
from app.models.scholarship import (
    ScholarshipStatus,
    ScholarshipSubTypeConfig,
    ScholarshipType,
    SubTypeSelectionMode,
)
from app.models.user import User, UserRole, UserType

URL = "/api/v1/admin/scholarships/{code}/college-applicant-stats/export"


@pytest_asyncio.fixture
async def phd(db: AsyncSession) -> ScholarshipType:
    st = ScholarshipType(
        code="cas_phd",
        name="博士生獎學金",
        status=ScholarshipStatus.active.value,
        sub_type_selection_mode=SubTypeSelectionMode.multiple,
    )
    db.add(st)
    await db.commit()
    await db.refresh(st)
    for order, (code, name) in enumerate([("nstc", "國科會博士生獎學金"), ("moe_1w", "教育部博士生獎學金")], 1):
        db.add(
            ScholarshipSubTypeConfig(
                scholarship_type_id=st.id, sub_type_code=code, name=name, display_order=order, is_active=True
            )
        )
    await db.commit()
    return st


async def _user(db: AsyncSession, nycu_id: str, role: UserRole) -> User:
    user = User(
        nycu_id=nycu_id,
        name=nycu_id,
        email=f"{nycu_id}@university.edu",
        user_type=UserType.student if role == UserRole.student else UserType.employee,
        role=role,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


@pytest_asyncio.fixture
async def as_user(client: AsyncClient):
    def _login(user: User) -> AsyncClient:
        app.dependency_overrides[get_current_user] = lambda: user
        return client

    yield _login
    app.dependency_overrides.pop(get_current_user, None)


async def _seed_applications(db: AsyncSession, st: ScholarshipType) -> None:
    rows = [
        # (student, academy, sub-types, status, year, semester)
        ("s1", "C", ["nstc"], "submitted", 115, None),
        ("s2", "C", ["nstc", "moe_1w"], "approved", 115, None),
        ("s3", "E", ["moe_1w"], "rejected", 115, None),
        ("s4", "E", ["nstc"], "draft", 115, None),  # never submitted ⇒ excluded
        ("s5", "E", ["nstc"], "deleted", 115, None),  # deleted ⇒ excluded
        ("s6", "E", ["nstc"], "submitted", 114, None),  # other year ⇒ excluded
        ("s7", "I", ["nstc"], "submitted", 115, "first"),
        ("s8", "S", ["nstc"], "approved", 115, "first"),  # s8 applied to each agency
        ("s8", "S", ["moe_1w"], "submitted", 115, "second"),  # separately ⇒ 國科會+教育部
    ]
    students = {}
    for index, (nycu_id, academy, sub_types, status, year, semester) in enumerate(rows):
        if nycu_id not in students:
            students[nycu_id] = await _user(db, f"cas_{nycu_id}", UserRole.student)
        student = students[nycu_id]
        db.add(
            Application(
                app_id=f"APP-CAS-{index}",
                user_id=student.id,
                scholarship_type_id=st.id,
                academic_year=year,
                semester=semester,
                status=status,
                scholarship_subtype_list=sub_types,
                sub_type_selection_mode=SubTypeSelectionMode.multiple,
                student_data={"std_stdcode": nycu_id, "std_academyno": academy},
            )
        )
    await db.commit()


def _sheet_counts(payload: bytes) -> dict:
    ws = load_workbook(io.BytesIO(payload)).active
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    counts = {
        ws.cell(row=r, column=2).value: tuple(ws.cell(row=r, column=c).value for c in range(3, ws.max_column + 1))
        for r in range(3, ws.max_row + 1)
    }
    return {"title": ws["A1"].value, "headers": headers, "counts": counts}


@pytest.mark.asyncio
async def test_admin_exports_counts_per_college_and_agency(db, phd, as_user):
    await _seed_applications(db, phd)
    admin = await _user(db, "cas_admin", UserRole.admin)

    resp = await as_user(admin).get(URL.format(code=phd.code), params={"academic_year": 115})

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert unquote(resp.headers["content-disposition"]).endswith("115學年度博士生獎學金_統計各學院學生申請人數.xlsx")
    sheet = _sheet_counts(resp.content)
    assert sheet["title"] == "115學年度博士生獎學金 統計各學院學生申請人數"
    assert sheet["headers"] == ["編號", "學院", "國科會", "教育部", "國科會+教育部"]
    # s2 applied to both agencies ⇒ only in 國科會+教育部
    assert sheet["counts"]["資訊學院"] == (1, 0, 1)
    assert sheet["counts"]["電機學院"] == (0, 1, 0)
    assert sheet["counts"]["工學院"] == (1, 0, 0)
    assert sheet["counts"]["理學院"] == (0, 0, 1)
    assert sheet["counts"]["合計"] == (2, 1, 2)


@pytest.mark.asyncio
async def test_semester_filter_limits_to_that_semester(db, phd, as_user):
    await _seed_applications(db, phd)
    admin = await _user(db, "cas_admin2", UserRole.admin)

    resp = await as_user(admin).get(URL.format(code=phd.code), params={"academic_year": 115, "semester": "first"})

    assert resp.status_code == 200, resp.text
    sheet = _sheet_counts(resp.content)
    assert sheet["title"].startswith("115學年度第一學期")
    assert sheet["counts"]["合計"] == (2, 0, 0)
    assert sheet["counts"]["工學院"] == (1, 0, 0)
    assert sheet["counts"]["理學院"] == (1, 0, 0)  # only s8's 國科會 application is in 第一學期


@pytest.mark.asyncio
async def test_non_admin_is_forbidden(db, phd, as_user):
    college = await _user(db, "cas_college", UserRole.college)
    resp = await as_user(college).get(URL.format(code=phd.code), params={"academic_year": 115})
    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_unknown_scholarship_is_404(db, as_user):
    admin = await _user(db, "cas_admin3", UserRole.admin)
    resp = await as_user(admin).get(URL.format(code="no_such_type"), params={"academic_year": 115})
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_invalid_semester_is_rejected(db, phd, as_user):
    admin = await _user(db, "cas_admin4", UserRole.admin)
    resp = await as_user(admin).get(URL.format(code=phd.code), params={"academic_year": 115, "semester": "x"})
    assert resp.status_code == 422
