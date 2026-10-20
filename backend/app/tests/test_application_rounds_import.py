"""申請梯次 column of the shared 批次匯入 / 補充匯入 workbook (多梯次申請 follow-up).

Rule for non-renewal rows: a single-round configuration takes a blank cell as
第一梯次; a configuration with later rounds requires the cell. The round must exist
and open the row's sub-types. The created application carries that round.
"""

import io
from datetime import datetime, timedelta, timezone
from unittest.mock import AsyncMock

import pandas as pd
import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.exceptions import ValidationError
from app.models.application import Application
from app.models.application_round import ScholarshipApplicationRound
from app.models.scholarship import ScholarshipConfiguration, ScholarshipType, SubTypeSelectionMode
from app.models.user import User, UserRole, UserType
from app.services.batch_import_service import BatchImportService
from app.services.batch_import_template_service import build_batch_import_template
from app.services.supplementary_import_service import SupplementaryImportService

NOW = datetime.now(timezone.utc)
PAST = (NOW - timedelta(days=40), NOW - timedelta(days=30))
OPEN = (NOW - timedelta(days=1), NOW + timedelta(days=5))
YEAR = 114


@pytest_asyncio.fixture
async def scholarship(db):
    scholarship = ScholarshipType(
        code="round_import_test",
        name="梯次匯入測試",
        sub_type_list=["nstc", "moe_1w"],
        sub_type_selection_mode=SubTypeSelectionMode.single,
        status="active",
    )
    db.add(scholarship)
    await db.flush()
    return scholarship


@pytest_asyncio.fixture
async def config(db, scholarship):
    config = ScholarshipConfiguration(
        scholarship_type_id=scholarship.id,
        academic_year=YEAR,
        semester=None,
        config_name="梯次匯入 114",
        config_code="round_import_114",
        amount=30000,
        application_start_date=PAST[0],
        application_end_date=PAST[1],
    )
    db.add(config)
    await db.flush()
    await db.refresh(config, attribute_names=["scholarship_type", "application_rounds"])
    return config


async def _open_second_round(db, config, sub_types=("nstc",)):
    db.add(
        ScholarshipApplicationRound(
            configuration_id=config.id,
            round_number=2,
            application_start_date=OPEN[0],
            application_end_date=OPEN[1],
            open_sub_types=list(sub_types),
        )
    )
    await db.flush()
    await db.refresh(config, attribute_names=["application_rounds"])


def _workbook(rows) -> bytes:
    """rows: (學號, 國科會, 教育部, 申請梯次 or None)."""
    df = pd.DataFrame(
        [
            {"學號": sid, "學生姓名": f"學生{sid}", "國科會": nstc, "教育部": moe, "申請梯次": round_cell}
            for sid, nstc, moe, round_cell in rows
        ]
    )
    buffer = io.BytesIO()
    df.to_excel(buffer, index=False)
    return buffer.getvalue()


async def _parse(db, scholarship, rows):
    return await BatchImportService(db).parse_excel_file(_workbook(rows), scholarship.id, YEAR, "yearly")


def _messages(errors) -> dict:
    return {error.student_id: error.message for error in errors}


class TestParser:
    async def test_single_round_configuration_accepts_a_blank_cell(self, db, scholarship, config):
        parsed, errors = await _parse(db, scholarship, [("310000001", 1, 0, None)])
        assert errors == []
        assert [row["student_id"] for row in parsed] == ["310000001"]

    async def test_multi_round_configuration_checks_every_row(self, db, scholarship, config):
        await _open_second_round(db, config, ("nstc",))
        parsed, errors = await _parse(
            db,
            scholarship,
            [
                ("310000011", 1, 0, 2),  # ok: round 2 opens nstc
                ("310000012", 1, 0, None),  # blank with two rounds
                ("310000013", 1, 0, 3),  # no such round
                ("310000014", 0, 1, 2),  # moe_1w closed in round 2
                ("310000015", 1, 0, "abc"),  # not a number
                ("310000016", 0, 1, 1),  # ok: round 1 opens every sub-type
            ],
        )
        assert sorted(row["student_id"] for row in parsed) == ["310000011", "310000016"]
        messages = _messages(errors)
        assert "請填寫「申請梯次」（可填：1、2）" in messages["310000012"]
        assert "申請梯次 3 不存在" in messages["310000013"]
        assert "第二梯次僅開放申請：nstc" in messages["310000014"]
        assert "申請梯次請填正整數" in messages["310000015"]

    async def test_round_one_restriction_applies_to_imports(self, db, scholarship, config):
        config.first_round_open_sub_types = ["nstc"]
        await db.flush()
        _, errors = await _parse(db, scholarship, [("310000021", 0, 1, None)])
        assert "第一梯次僅開放申請：nstc" in _messages(errors)["310000021"]


def _row(student_id: str, sub_types=("nstc",), application_round=None):
    return {
        "student_id": student_id,
        "student_name": "測試生",
        "postal_account": None,
        "advisor_name": None,
        "advisor_email": None,
        "advisor_nycu_id": None,
        "is_renewal": False,
        "renewal_year": None,
        "application_round": application_round,
        "sub_types": list(sub_types),
        "custom_fields": {},
    }


async def _importer(db):
    importer = User(
        nycu_id="round_importer",
        name="College",
        email="round_importer@nycu.edu.tw",
        user_type=UserType.employee,
        role=UserRole.college,
        college_code="A",
    )
    db.add(importer)
    await db.flush()
    return importer


class TestSupplementaryImportCreatesTheRound:
    async def _create(self, db, config, rows):
        service = SupplementaryImportService(db, student_service=AsyncMock())
        data_map = {row["student_id"]: {"std_stdcode": row["student_id"], "std_cname": "新生"} for row in rows}
        user_map = await service.find_or_create_users(data_map)
        importer = await _importer(db)
        await service.create_applications(rows, user_map, data_map, config, importer_id=importer.id)
        await db.flush()
        result = await db.execute(select(User.nycu_id, Application.application_round).join(Application.student))
        return dict(result.all())

    async def test_application_gets_the_row_round_not_the_open_round(self, db, config):
        await _open_second_round(db, config, ("nstc", "moe_1w"))
        rounds = await self._create(
            db, config, [_row("310000031", application_round=1), _row("310000032", application_round=2)]
        )
        assert rounds == {"310000031": 1, "310000032": 2}

    async def test_single_round_blank_cell_is_round_one(self, db, config):
        assert await self._create(db, config, [_row("310000041")]) == {"310000041": 1}

    async def test_closed_sub_type_is_refused_with_the_student_id(self, db, config):
        await _open_second_round(db, config, ("nstc",))
        with pytest.raises(ValidationError, match="學號 310000051：第二梯次僅開放申請"):
            await self._create(db, config, [_row("310000051", ("moe_1w",), application_round=2)])


async def test_template_has_the_application_round_column(db, scholarship):
    worksheet = load_workbook(io.BytesIO(await build_batch_import_template(db, scholarship))).active
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    assert "申請梯次" in headers
    column = headers.index("申請梯次") + 1
    assert "1 = 第一梯次" in worksheet.cell(row=1, column=column).comment.text
    assert worksheet.cell(row=2, column=column).value == 1
