"""指導教授審核結果/評語 in the exports.

- professor_review_export: loading professor verdicts (professors only) and
  the shared cell wording (推薦 / 不推薦 / 未審核 / —, co-advisor joins);
- CollegeRankingExportService: the result columns, then the comment columns,
  right after 指導教授姓名, in both xlsx and PDF;
- professor_review_pdf: the 合併檔's 指導教授評語 page.
"""

import io
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.application import Application
from app.models.enums import ApplicationStatus, Semester
from app.models.review import ApplicationReview, ApplicationReviewItem
from app.models.scholarship import ScholarshipType
from app.models.user import User, UserRole, UserType
from app.services.college_ranking_export_service import STATIC_HEADERS, CollegeRankingExportService, ExportRow
from app.services.professor_review_export import (
    DEFAULT_SUB_TYPE,
    NOT_APPLIED_LABEL,
    NOT_REVIEWED_LABEL,
    ProfessorReviewEntry,
    applied_sub_types,
    load_professor_reviews,
    review_cell_pair,
    review_cells,
    review_headers,
    review_sub_type_columns,
)
from app.services.professor_review_pdf import _rows, build_professor_review_pdf

LABELS = {"nstc": "國科會", "moe_1w": "教育部(一萬)"}
REVIEWED_AT = datetime(2026, 10, 5, 6, 30, tzinfo=timezone.utc)  # 14:30 Taipei


def _entry(code, rec, comments=None, name="王教授"):
    return ProfessorReviewEntry(
        sub_type_code=code, recommendation=rec, comments=comments, reviewer_name=name, reviewed_at=REVIEWED_AT
    )


def _app(sub_types=("nstc", "moe_1w")):
    return SimpleNamespace(
        scholarship_subtype_list=list(sub_types),
        sub_type_preferences=None,
        sub_scholarship_type=None,
        student_data={"std_stdcode": "001", "std_cname": "甲"},
        submitted_form_data={},
    )


# ─── column set ──────────────────────────────────────────────────────


def test_columns_follow_the_scholarship_sub_type_order():
    stype = SimpleNamespace(sub_type_list=["NSTC ", "moe_1w", "nstc"])
    assert review_sub_type_columns(stype) == ["nstc", "moe_1w"]


@pytest.mark.parametrize("sub_type_list", [None, [], ["general"]])
def test_scholarship_without_sub_types_gets_one_default_pair(sub_type_list):
    stype = SimpleNamespace(sub_type_list=sub_type_list)
    assert review_sub_type_columns(stype) == [DEFAULT_SUB_TYPE]
    assert review_headers([DEFAULT_SUB_TYPE], LABELS) == ["指導教授審核結果", "指導教授評語"]


def test_headers_put_every_result_before_every_comment():
    assert review_headers(["nstc", "moe_1w"], LABELS) == [
        "國科會-指導教授審核結果",
        "教育部(一萬)-指導教授審核結果",
        "國科會-指導教授評語",
        "教育部(一萬)-指導教授評語",
    ]


def test_applied_sub_types_falls_back_to_default_like_review_service():
    assert applied_sub_types(_app(())) == [DEFAULT_SUB_TYPE]
    assert applied_sub_types(_app((" NSTC",))) == ["nstc"]


# ─── cell wording ────────────────────────────────────────────────────


def test_cells_render_verdict_comment_and_missing_states():
    app = _app(("nstc",))
    entries = [_entry("nstc", "reject", "已獲其他補助")]
    assert review_cell_pair(app, entries, "nstc") == ("不推薦", "已獲其他補助")
    # Applied for moe_1w? No → nothing for the professor to review.
    assert review_cell_pair(app, entries, "moe_1w") == (NOT_APPLIED_LABEL, "")
    # Applied, no verdict yet
    assert review_cell_pair(app, [], "nstc") == (NOT_REVIEWED_LABEL, "")


def test_cells_order_matches_headers():
    app = _app()
    entries = [_entry("nstc", "approve", "研究表現優異"), _entry("moe_1w", "reject", None)]
    assert review_cells(app, entries, ["nstc", "moe_1w"]) == ["推薦", "不推薦", "研究表現優異", ""]


def test_co_advisors_on_one_sub_type_are_all_kept():
    app = _app(("nstc",))
    entries = [_entry("nstc", "approve", "很好", name="王教授"), _entry("nstc", "reject", None, name="李教授")]
    result, comment = review_cell_pair(app, entries, "nstc")
    assert result == "王教授：推薦；李教授：不推薦"
    assert comment == "王教授：很好"


# ─── 學生資料彙整表 columns ──────────────────────────────────────────


def _row(entries):
    return ExportRow(rank_position=1, application=_app(), advisor_names="王教授", professor_reviews=entries)


def test_workbook_places_review_columns_after_advisor_name():
    service = CollegeRankingExportService()
    payload = service.build_workbook(
        rows=[_row([_entry("nstc", "approve", "研究表現優異"), _entry("moe_1w", "reject", "已獲國科會支持")])],
        dynamic_fields=[],
        sub_type_labels=LABELS,
        title="t",
        sheet_name="s",
        review_sub_types=["nstc", "moe_1w"],
    )
    ws = load_workbook(io.BytesIO(payload)).active
    headers = [c.value for c in ws[2]]
    advisor_idx = headers.index("指導教授姓名")
    assert advisor_idx == len(STATIC_HEADERS) - 1
    assert headers[advisor_idx + 1 : advisor_idx + 5] == [
        "國科會-指導教授審核結果",
        "教育部(一萬)-指導教授審核結果",
        "國科會-指導教授評語",
        "教育部(一萬)-指導教授評語",
    ]
    values = [c.value for c in ws[3]]
    assert values[advisor_idx + 1 : advisor_idx + 5] == ["推薦", "不推薦", "研究表現優異", "已獲國科會支持"]


def test_workbook_without_review_sub_types_has_no_review_columns():
    payload = CollegeRankingExportService().build_workbook(
        rows=[_row([])], dynamic_fields=[], sub_type_labels=LABELS, title="t", sheet_name="s"
    )
    headers = [c.value for c in load_workbook(io.BytesIO(payload)).active[2]]
    assert headers == STATIC_HEADERS


def test_pdf_widths_cover_review_and_seal_columns():
    service = CollegeRankingExportService()
    headers = service._headers([], LABELS, ["nstc", "moe_1w"], include_college_seal=True)
    widths = service._pdf_col_widths(785.0, 0, 2, include_college_seal=True)
    assert len(widths) == len(headers)
    assert sum(widths) == pytest.approx(785.0)
    comment_idx = headers.index("國科會-指導教授評語")
    assert widths[comment_idx] > widths[comment_idx - 2]  # comment wider than its verdict


def test_pdf_renders_long_review_comment():
    out = CollegeRankingExportService().build_pdf(
        rows=[_row([_entry("nstc", "approve", "研究表現優異" * 400)])],
        dynamic_fields=[],
        sub_type_labels=LABELS,
        title="t",
        include_college_seal=True,
        review_sub_types=["nstc", "moe_1w"],
    )
    assert out[:5] == b"%PDF-"


# ─── 合併檔 指導教授評語 page ────────────────────────────────────────


def test_pdf_page_rows_cover_every_column_and_stray_verdicts():
    app = _app(("nstc",))
    entries = [_entry("nstc", "approve", "很好"), _entry(DEFAULT_SUB_TYPE, "reject", "舊資料")]
    rows = _rows(app, entries, ["nstc", "moe_1w"], LABELS)
    assert rows == [
        ["國科會", "推薦", "很好", "王教授", "2026-10-05 14:30"],
        ["教育部(一萬)", NOT_APPLIED_LABEL, "", "", ""],
        ["—", "不推薦", "舊資料", "王教授", "2026-10-05 14:30"],
    ]


def test_pdf_page_renders_text_and_survives_a_page_long_comment():
    out = build_professor_review_pdf(
        subtitle_lines=["某獎學金 114學年度 全學年", "001 甲"],
        application=_app(),
        entries=[_entry("nstc", "approve", "研究表現<優異>\n第二行"), _entry("moe_1w", "reject", "長" * 6000)],
        review_sub_types=["nstc", "moe_1w"],
        sub_type_labels=LABELS,
    )
    reader = PdfReader(io.BytesIO(out))
    assert len(reader.pages) >= 2  # the long comment flowed onto another page
    text = reader.pages[0].extract_text()
    assert "指導教授評語" in text
    assert "國科會" in text
    assert "研究表現<優異>" in text
    assert "王教授" in text


# ─── loader ──────────────────────────────────────────────────────────


async def _user(db: AsyncSession, suffix: str, role: UserRole, name: str) -> User:
    user = User(
        nycu_id=f"pre_{suffix}",
        name=name,
        email=f"pre_{suffix}@u.edu",
        user_type=UserType.student if role == UserRole.student else UserType.employee,
        role=role,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


async def _review(db: AsyncSession, app_id: int, reviewer: User, items, reviewed_at: datetime) -> None:
    review = ApplicationReview(
        application_id=app_id, reviewer_id=reviewer.id, recommendation="approve", reviewed_at=reviewed_at
    )
    db.add(review)
    await db.commit()
    await db.refresh(review)
    for code, rec, comments in items:
        db.add(ApplicationReviewItem(review_id=review.id, sub_type_code=code, recommendation=rec, comments=comments))
    await db.commit()


@pytest.mark.asyncio
async def test_loader_returns_only_professor_verdicts_with_reviewer(db: AsyncSession):
    student = await _user(db, "stu", UserRole.student, "學生")
    prof = await _user(db, "prof", UserRole.professor, "王教授")
    admin = await _user(db, "adm", UserRole.admin, "管理員")
    sch = ScholarshipType(code="pre_sch", name="Pre Sch", status="active")
    db.add(sch)
    await db.commit()
    await db.refresh(sch)
    app = Application(
        app_id="APP-PRE-1",
        user_id=student.id,
        scholarship_type_id=sch.id,
        sub_type_selection_mode="multiple",
        academic_year=114,
        semester=Semester.first.value,
        status=ApplicationStatus.submitted.value,
        scholarship_subtype_list=["nstc", "moe_1w"],
        agree_terms=True,
    )
    db.add(app)
    await db.commit()
    await db.refresh(app)

    await _review(db, app.id, prof, [(" NSTC", "approve", "很好"), ("moe_1w", "reject", None)], REVIEWED_AT)
    await _review(db, app.id, admin, [("nstc", "reject", "admin note")], REVIEWED_AT)

    result = await load_professor_reviews(db, [app.id])

    assert [(e.sub_type_code, e.recommendation, e.comments, e.reviewer_name) for e in result[app.id]] == [
        ("nstc", "approve", "很好", "王教授"),
        ("moe_1w", "reject", None, "王教授"),
    ]
    assert await load_professor_reviews(db, []) == {}
