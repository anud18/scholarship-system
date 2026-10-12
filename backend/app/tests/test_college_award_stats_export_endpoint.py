"""Endpoint tests for the 統計各學院學生獲獎人數 exports.

- GET /api/v1/manual-distribution/distribution-summary/college-award-stats/export
  (current 分發名單)
- GET /api/v1/payment-rosters/{roster_id}/college-award-stats/export
  (one 造冊's included students)

Counts are asserted by parsing the workbook, never by scanning response bytes.
"""

import io
from datetime import datetime, timezone
from urllib.parse import unquote

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_admin_user, get_current_user
from app.main import app
from app.models.application import Application
from app.models.college_review import CollegeRanking, CollegeRankingItem
from app.models.enums import ApplicationStatus, ReviewStage, SubTypeSelectionMode
from app.models.payment_roster import (
    PaymentRoster,
    PaymentRosterItem,
    RosterCycle,
    RosterStatus,
    RosterTriggerType,
)
from app.models.scholarship import ScholarshipConfiguration, ScholarshipSubTypeConfig, ScholarshipType
from app.models.user import User, UserRole, UserType
from app.services.college_stats_export_service import AWARD_REPORT_NAME, TOTAL_ROW_LABEL

ACADEMIC_YEAR = 114
DISTRIBUTION_URL = "/api/v1/manual-distribution/distribution-summary/college-award-stats/export"
ROSTER_URL = "/api/v1/payment-rosters/{roster_id}/college-award-stats/export"


async def _user(db: AsyncSession, nycu_id: str, role: UserRole) -> User:
    user = User(
        nycu_id=nycu_id,
        name=nycu_id,
        email=f"{nycu_id}@university.edu",
        user_type=UserType.student if role == UserRole.student else UserType.employee,
        role=role,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


@pytest_asyncio.fixture
async def admin(db: AsyncSession) -> User:
    return await _user(db, "award_admin", UserRole.admin)


@pytest_asyncio.fixture
async def admin_client(client: AsyncClient, admin: User):
    app.dependency_overrides[get_current_admin_user] = lambda: admin
    app.dependency_overrides[get_current_user] = lambda: admin
    try:
        yield client
    finally:
        app.dependency_overrides.pop(get_current_admin_user, None)
        app.dependency_overrides.pop(get_current_user, None)


@pytest_asyncio.fixture
async def phd(db: AsyncSession):
    st = ScholarshipType(code="award_phd", name="博士生獎學金")
    db.add(st)
    await db.commit()
    await db.refresh(st)
    config = ScholarshipConfiguration(
        scholarship_type_id=st.id,
        academic_year=ACADEMIC_YEAR,
        semester=None,
        config_name="Award Config",
        config_code="award-config",
        amount=40000,
        currency="TWD",
        is_active=True,
    )
    db.add(config)
    for order, (code, name) in enumerate(
        [("nstc", "國科會博士生獎學金"), ("moe_1w", "教育部 1 萬"), ("moe_2w", "教育部 2 萬")], 1
    ):
        db.add(
            ScholarshipSubTypeConfig(
                scholarship_type_id=st.id, sub_type_code=code, name=name, display_order=order, is_active=True
            )
        )
    await db.commit()
    await db.refresh(config)
    return st, config


async def _application(
    db: AsyncSession, st: ScholarshipType, suffix: str, academy: str, sub_type: str, trm_academy: str | None = None
) -> Application:
    user = await _user(db, f"award_stu_{suffix}", UserRole.student)
    application = Application(
        app_id=f"APP-{ACADEMIC_YEAR}-0-8{suffix}",
        user_id=user.id,
        scholarship_type_id=st.id,
        scholarship_subtype_list=[sub_type],
        sub_type_selection_mode=SubTypeSelectionMode.single,
        sub_scholarship_type=sub_type,
        academic_year=ACADEMIC_YEAR,
        semester=None,
        status=ApplicationStatus.approved,
        review_stage=ReviewStage.quota_distributed,
        agree_terms=True,
        student_data={
            "std_stdcode": f"31{suffix}",
            "std_cname": suffix,
            "std_academyno": academy,
            **({"trm_academyno": trm_academy} if trm_academy else {}),
        },
    )
    db.add(application)
    await db.commit()
    await db.refresh(application)
    return application


def _sheet_counts(content: bytes) -> tuple:
    ws = load_workbook(io.BytesIO(content)).active
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    counts = {
        ws.cell(row=r, column=2).value: tuple(ws.cell(row=r, column=c).value for c in range(3, ws.max_column + 1))
        for r in range(3, ws.max_row + 1)
    }
    return ws, headers, counts


async def _seed_distribution(
    db: AsyncSession, st: ScholarshipType, *, executed: bool = True, allocate: bool = True
) -> None:
    ranking = CollegeRanking(
        scholarship_type_id=st.id,
        sub_type_code="default",
        academic_year=ACADEMIC_YEAR,
        semester=None,
        is_finalized=True,
        distribution_executed=executed,
        ranking_status="finalized",
    )
    db.add(ranking)
    await db.commit()
    await db.refresh(ranking)
    rows = [
        # (student, academy, sub-type, allocated?)
        ("d1", "C", "nstc", True),  # trm_academyno E: panel buckets by std first ⇒ 資訊
        ("d2", "C", "moe_1w", True),
        ("d3", "E", "moe_2w", True),
        ("d4", "E", "nstc", False),  # ranked but not allocated ⇒ not an award
    ]
    for rank, (suffix, academy, sub_type, allocated) in enumerate(rows, 1):
        allocated = allocated and allocate
        application = await _application(db, st, suffix, academy, sub_type, trm_academy="E" if suffix == "d1" else None)
        db.add(
            CollegeRankingItem(
                ranking_id=ranking.id,
                application_id=application.id,
                rank_position=rank,
                is_allocated=allocated,
                allocated_sub_type=sub_type if allocated else None,
                status="allocated" if allocated else "ranked",
            )
        )
    await db.commit()


@pytest.mark.asyncio
async def test_distribution_award_stats_counts_allocated_students(admin_client, db, phd):
    st, _ = phd
    await _seed_distribution(db, st)

    resp = await admin_client.get(
        DISTRIBUTION_URL,
        params={"scholarship_type_id": st.id, "academic_year": ACADEMIC_YEAR, "semester": "yearly"},
    )

    assert resp.status_code == 200, resp.text
    assert AWARD_REPORT_NAME in unquote(resp.headers["content-disposition"])
    ws, headers, counts = _sheet_counts(resp.content)
    assert ws.title == AWARD_REPORT_NAME
    assert ws["A1"].value == f"{ACADEMIC_YEAR}學年度博士生獎學金 {AWARD_REPORT_NAME}"
    assert headers == ["編號", "學院", "國科會", "教育部"]
    assert counts["資訊學院"] == (1, 1)
    assert counts["電機學院"] == (0, 1)
    assert counts[TOTAL_ROW_LABEL] == (1, 2)


@pytest.mark.asyncio
async def test_distribution_award_stats_404_before_distribution_runs(admin_client, db, phd):
    st, _ = phd
    await _seed_distribution(db, st, executed=False)

    resp = await admin_client.get(
        DISTRIBUTION_URL,
        params={"scholarship_type_id": st.id, "academic_year": ACADEMIC_YEAR, "semester": "yearly"},
    )

    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_distribution_award_stats_404_when_nobody_is_allocated(admin_client, db, phd):
    st, _ = phd
    await _seed_distribution(db, st, allocate=False)

    resp = await admin_client.get(
        DISTRIBUTION_URL,
        params={"scholarship_type_id": st.id, "academic_year": ACADEMIC_YEAR, "semester": "yearly"},
    )

    assert resp.status_code == 404
    assert "尚無已分配" in resp.json()["message"]


@pytest.mark.asyncio
async def test_distribution_award_stats_rejects_unknown_semester(admin_client, phd):
    st, _ = phd
    resp = await admin_client.get(
        DISTRIBUTION_URL,
        params={"scholarship_type_id": st.id, "academic_year": ACADEMIC_YEAR, "semester": "summer"},
    )
    assert resp.status_code == 422


async def _seed_roster(
    db: AsyncSession,
    st: ScholarshipType,
    config: ScholarshipConfiguration,
    creator: User,
    roster_status: RosterStatus = RosterStatus.COMPLETED,
):
    roster = PaymentRoster(
        roster_code="ROSTER-AWARD-2026-01",
        scholarship_configuration_id=config.id,
        period_label="2026-01",
        academic_year=ACADEMIC_YEAR,
        roster_cycle=RosterCycle.MONTHLY,
        status=roster_status,
        trigger_type=RosterTriggerType.MANUAL,
        created_by=creator.id,
        sub_type="nstc",
        started_at=datetime.now(timezone.utc),
        completed_at=datetime.now(timezone.utc),
    )
    db.add(roster)
    await db.commit()
    await db.refresh(roster)
    rows = [
        # (student, academy, allocated_sub_type snapshot, included?)
        ("r1", "C", "nstc", True),
        ("r2", "I", "nstc", True),
        ("r3", "I", None, True),  # no item snapshot ⇒ the roster's sub_type, not the applied-for one
        ("r4", "C", "nstc", False),  # excluded / removed ⇒ not counted
    ]
    for suffix, academy, sub_type, included in rows:
        # Applied for moe_1w: proves the count never reads the applied-for sub-type.
        application = await _application(db, st, suffix, academy, "moe_1w")
        db.add(
            PaymentRosterItem(
                roster_id=roster.id,
                application_id=application.id,
                student_id_number=f"A1{suffix}",
                student_name=suffix,
                scholarship_name=st.name,
                scholarship_amount=40000,
                allocated_sub_type=sub_type,
                scholarship_subtype="moe_1w",
                is_included=included,
            )
        )
    await db.commit()
    return roster


@pytest.mark.asyncio
async def test_roster_award_stats_counts_included_items(admin_client, db, phd, admin):
    st, config = phd
    roster = await _seed_roster(db, st, config, admin)

    resp = await admin_client.get(ROSTER_URL.format(roster_id=roster.id))

    assert resp.status_code == 200, resp.text
    assert unquote(resp.headers["content-disposition"]).endswith(f"ROSTER-AWARD-2026-01_{AWARD_REPORT_NAME}.xlsx")
    ws, headers, counts = _sheet_counts(resp.content)
    assert "2026-01" in ws["A1"].value
    assert headers == ["編號", "學院", "國科會", "教育部"]
    assert counts["資訊學院"] == (1, 0)
    assert counts["工學院"] == (2, 0)
    assert counts[TOTAL_ROW_LABEL] == (3, 0)


@pytest.mark.asyncio
async def test_roster_award_stats_409_for_unfinished_roster(admin_client, db, phd, admin):
    st, config = phd
    roster = await _seed_roster(db, st, config, admin, roster_status=RosterStatus.FAILED)

    resp = await admin_client.get(ROSTER_URL.format(roster_id=roster.id))

    assert resp.status_code == 409


@pytest.mark.asyncio
async def test_roster_award_stats_404_for_missing_roster(admin_client):
    resp = await admin_client.get(ROSTER_URL.format(roster_id=999999))
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_roster_award_stats_forbidden_for_non_admin(client, db, phd, admin):
    st, config = phd
    roster = await _seed_roster(db, st, config, admin)
    student = await _user(db, "award_nosy", UserRole.student)
    app.dependency_overrides[get_current_user] = lambda: student
    try:
        resp = await client.get(ROSTER_URL.format(roster_id=roster.id))
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert resp.status_code == 403
