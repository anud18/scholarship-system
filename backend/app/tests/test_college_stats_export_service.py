"""Pure-logic tests for the 統計各學院學生申請人數 / 獲獎人數 export service."""

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.models.enums import SubTypeSelectionMode

from app.services.college_stats_export_service import (
    APPLICANT_COUNT_LABEL,
    AWARD_COUNT_LABEL,
    AWARD_REPORT_NAME,
    COLLEGE_STATS_LAYOUT,
    SINGLE_COLUMN_KEY,
    TOTAL_ROW_LABEL,
    UNKNOWN_COLLEGE_LABEL,
    UNSPECIFIED_SUB_TYPE_COLUMN,
    StatsRecord,
    StatsColumn,
    build_stats_columns,
    build_workbook,
    scholarship_stats_columns,
    tally_college_counts,
)

PHD_CONFIGS = [("nstc", "國科會博士生獎學金"), ("moe_1w", "教育部 1 萬"), ("moe_2w", "教育部 2 萬")]


def _rec(user_id, academy, *sub_types):
    return StatsRecord(user_id=user_id, academy_code=academy, sub_type_codes=tuple(sub_types))


def _counts(rows):
    return {row.label: row.counts for row in rows}


def test_moe_sub_types_collapse_into_one_agency_column():
    columns = build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "nstc")])
    assert [c.label for c in columns] == ["國科會", "教育部"]


def test_ungrouped_sub_type_uses_its_configured_name():
    columns = build_stats_columns([("nstc", "x"), ("custom_new", "自訂獎助")], [])
    assert [c.label for c in columns] == ["國科會", "自訂獎助"]


def test_unconfigured_sub_type_on_an_application_still_gets_a_column():
    columns = build_stats_columns([("nstc", "x")], [_rec(1, "C", "legacy_type")])
    assert [c.key for c in columns] == ["國科會", "legacy_type"]


def test_scholarship_without_sub_types_gets_single_applicant_column():
    columns = build_stats_columns([], [_rec(1, "C", "general")])
    assert columns == [StatsColumn(SINGLE_COLUMN_KEY, APPLICANT_COUNT_LABEL)]
    rows = tally_college_counts([_rec(1, "C", "general"), _rec(2, "C")], columns)
    assert _counts(rows)["資訊學院"] == (2,)


def test_unspecified_column_only_when_an_application_lacks_a_sub_type():
    assert UNSPECIFIED_SUB_TYPE_COLUMN not in [c.key for c in build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "nstc")])]
    columns = build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "general")])
    assert columns[-1].key == UNSPECIFIED_SUB_TYPE_COLUMN
    assert _counts(tally_college_counts([_rec(1, "C", "general")], columns))["資訊學院"] == (0, 0, 1)


def test_tally_counts_each_student_once_per_column():
    records = [
        _rec(1, "C", "nstc", "moe_1w"),  # both agencies
        _rec(2, "C", "moe_1w", "moe_2w"),  # two MOE sub-types ⇒ one 教育部 count
        _rec(3, "E", "nstc"),
        _rec(3, "E", "nstc"),  # duplicate application, same student
    ]
    columns = build_stats_columns(PHD_CONFIGS, records)
    counts = _counts(tally_college_counts(records, columns))
    assert counts["資訊學院"] == (1, 2)
    assert counts["電機學院"] == (1, 0)
    assert counts[TOTAL_ROW_LABEL] == (2, 2)


def test_every_layout_college_is_listed_in_reference_order_even_at_zero():
    rows = tally_college_counts([], build_stats_columns(PHD_CONFIGS, []))
    assert [r.label for r in rows] == [label for label, _ in COLLEGE_STATS_LAYOUT] + [TOTAL_ROW_LABEL]
    assert all(r.counts == (0, 0) for r in rows)


def test_yangming_humanities_code_8_shares_the_humanities_row():
    records = [_rec(1, "A", "nstc"), _rec(2, "8", "nstc")]
    counts = _counts(tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records)))
    assert counts["人文藝術與社會學院"] == (2, 0)


def test_colleges_outside_layout_are_appended_before_total():
    records = [_rec(1, "X", "nstc"), _rec(2, None, "nstc"), _rec(3, "~", "nstc")]
    rows = tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records))
    tail = [r.label for r in rows[len(COLLEGE_STATS_LAYOUT) :]]
    assert tail == ["電機資訊學院", f"{UNKNOWN_COLLEGE_LABEL}（~）", UNKNOWN_COLLEGE_LABEL, TOTAL_ROW_LABEL]
    assert rows[-1].counts == (3, 0)


def test_retired_eecs_code_counts_in_ece_row():
    """電資學院 (Y) is merged into 電機學院: a stray Y snapshot adds no extra row."""
    records = [_rec(1, "E", "nstc"), _rec(2, "Y", "nstc")]
    rows = tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records))
    by_label = {r.label: r.counts for r in rows}
    assert by_label["電機學院"] == (2, 0)
    assert len(rows) == len(COLLEGE_STATS_LAYOUT) + 1


def test_workbook_matches_reference_layout():
    records = [_rec(1, "I", "nstc"), _rec(2, "I", "moe_1w")]
    columns = build_stats_columns(PHD_CONFIGS, records)
    payload = build_workbook(
        title="115學年度博士生獎學金 統計各學院學生申請人數",
        columns=columns,
        rows=tally_college_counts(records, columns),
        sheet_name="統計各學院學生申請人數",
    )

    ws = load_workbook(io.BytesIO(payload)).active
    assert ws["A1"].value == "115學年度博士生獎學金 統計各學院學生申請人數"
    assert "A1:D1" in {str(r) for r in ws.merged_cells.ranges}
    assert [ws.cell(row=2, column=c).value for c in range(1, 5)] == ["編號", "學院", "國科會", "教育部"]
    assert [ws.cell(row=4, column=c).value for c in range(1, 5)] == [2, "工學院", 1, 1]
    last = 2 + len(COLLEGE_STATS_LAYOUT) + 1
    assert [ws.cell(row=last, column=c).value for c in range(1, 5)] == [None, TOTAL_ROW_LABEL, 1, 1]
    assert ws["B3"].font.name == "微軟正黑體"


def test_workbook_neutralises_formula_like_sub_type_names():
    columns = build_stats_columns([("custom", "=HYPERLINK(1)")], [])
    payload = build_workbook(title="t", columns=columns, rows=tally_college_counts([], columns), sheet_name="s")
    ws = load_workbook(io.BytesIO(payload)).active
    assert not str(ws["C2"].value).startswith("=")


def test_total_row_is_the_column_sum_even_when_a_student_spans_two_colleges():
    records = [_rec(1, "C", "nstc"), _rec(1, "E", "nstc")]  # college changed between applications
    counts = _counts(tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records)))
    assert counts["資訊學院"] == (1, 0) and counts["電機學院"] == (1, 0)
    assert counts[TOTAL_ROW_LABEL] == (2, 0)


def test_deactivated_sub_type_column_keeps_its_configured_name():
    columns = build_stats_columns(
        [("nstc", "x")],
        [_rec(1, "C", "custom_grant")],
        known_names={"nstc": "x", "custom_grant": "自訂獎助"},
    )
    assert [c.label for c in columns] == ["國科會", "自訂獎助"]


def test_combined_agencies_column_counts_multi_agency_applicants_only_there():
    records = [
        _rec(1, "C", "nstc", "moe_1w"),  # both agencies ⇒ only 國科會+教育部
        _rec(2, "C", "moe_1w", "moe_2w"),  # two MOE sub-types ⇒ still just 教育部
        _rec(3, "C", "nstc"),
        _rec(4, "E", "moe_2w", "nstc"),
    ]
    columns = build_stats_columns(PHD_CONFIGS, records, combine_agencies=True)
    assert [c.label for c in columns] == ["國科會", "教育部", "國科會+教育部"]
    counts = _counts(tally_college_counts(records, columns, combine_agencies=True))
    assert counts["資訊學院"] == (1, 1, 1)
    assert counts["電機學院"] == (0, 0, 1)
    assert counts[TOTAL_ROW_LABEL] == (1, 1, 2)


def test_combined_column_is_on_the_form_even_when_nobody_picked_both():
    columns = build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "nstc")], combine_agencies=True)
    assert [c.label for c in columns] == ["國科會", "教育部", "國科會+教育部"]
    assert _counts(tally_college_counts([], columns, combine_agencies=True))[TOTAL_ROW_LABEL] == (0, 0, 0)


def test_no_combined_column_for_a_single_agency_scholarship():
    columns = build_stats_columns([("moe_1w", "a"), ("moe_2w", "b")], [_rec(1, "C", "moe_1w")], combine_agencies=True)
    assert [c.label for c in columns] == ["教育部"]


def test_multi_select_form_lists_every_combination_in_a_stable_order():
    configs = [("nstc", "x"), ("moe_1w", "y"), ("custom", "自訂獎助")]
    records = [_rec(1, "C", "custom", "nstc"), _rec(2, "C")]
    columns = build_stats_columns(configs, records, combine_agencies=True)
    assert [c.label for c in columns] == [
        "國科會",
        "教育部",
        "自訂獎助",
        "國科會+教育部",
        "國科會+自訂獎助",
        "教育部+自訂獎助",
        "國科會+教育部+自訂獎助",
        UNSPECIFIED_SUB_TYPE_COLUMN,
    ]
    assert columns[:-1] == build_stats_columns(configs, [], combine_agencies=True)
    assert _counts(tally_college_counts(records, columns, combine_agencies=True))["資訊學院"] == (
        0,
        0,
        0,
        0,
        1,
        0,
        0,
        1,
    )


def test_agencies_from_separate_applications_merge_per_student():
    records = [
        _rec(1, "C", "nstc"),  # e.g. a 續領 application …
        _rec(1, "C", "moe_1w"),  # … plus a new one ⇒ one student, 國科會+教育部
        _rec(2, "C", "nstc", "moe_1w"),
        _rec(2, "C", "nstc"),  # second application adds nothing new
        _rec(3, "C", "general"),
        _rec(3, "C", "moe_2w"),  # a real pick wins over a placeholder application
    ]
    columns = build_stats_columns(PHD_CONFIGS, records, combine_agencies=True)
    counts = _counts(tally_college_counts(records, columns, combine_agencies=True))
    assert [c.label for c in columns] == ["國科會", "教育部", "國科會+教育部"]
    assert counts["資訊學院"] == (0, 1, 2)
    assert sum(counts["資訊學院"]) == 3  # row sum = head count


def test_single_select_scholarship_gets_no_fixed_combination_column():
    columns = build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "nstc")], combine_agencies=True, allows_multiple=False)
    assert [c.label for c in columns] == ["國科會", "教育部"]


def test_tally_refuses_a_student_without_a_matching_column():
    columns = build_stats_columns(PHD_CONFIGS, [], combine_agencies=True, allows_multiple=False)
    with pytest.raises(ValueError):
        tally_college_counts([_rec(1, "C", "nstc", "moe_1w")], columns, combine_agencies=True)


def test_award_sheet_keeps_counting_each_agency_without_combinations():
    records = [_rec(1, "C", "nstc", "moe_1w")]
    columns = scholarship_stats_columns(_scholarship(PHD_CONFIGS), records, single_column_label=AWARD_COUNT_LABEL)
    assert [c.label for c in columns] == ["國科會", "教育部"]
    assert _counts(tally_college_counts(records, columns))["資訊學院"] == (1, 1)


def test_scholarship_columns_follow_the_selection_mode():
    records = [_rec(1, "C", "nstc")]
    multi = scholarship_stats_columns(
        _scholarship(PHD_CONFIGS), records, single_column_label=APPLICANT_COUNT_LABEL, combine_agencies=True
    )
    single = scholarship_stats_columns(
        _scholarship(PHD_CONFIGS, mode=SubTypeSelectionMode.single),
        records,
        single_column_label=APPLICANT_COUNT_LABEL,
        combine_agencies=True,
    )
    assert [c.label for c in multi] == ["國科會", "教育部", "國科會+教育部"]
    assert [c.label for c in single] == ["國科會", "教育部"]


def _scholarship(configs, inactive=(), mode=SubTypeSelectionMode.multiple):
    sub_type_configs = [
        SimpleNamespace(sub_type_code=code, name=name, is_active=True, display_order=i)
        for i, (code, name) in enumerate(configs, 1)
    ] + [SimpleNamespace(sub_type_code=code, name=name, is_active=False, display_order=99) for code, name in inactive]
    return SimpleNamespace(
        sub_type_configs=sub_type_configs,
        sub_type_selection_mode=mode,
        get_active_sub_type_configs=lambda: [c for c in sub_type_configs if c.is_active],
    )


def test_award_workbook_uses_award_sheet_name_and_agency_columns():
    records = [_rec(1, "C", "nstc"), _rec(2, "C", "moe_2w"), _rec(3, "E", "moe_1w")]
    columns = scholarship_stats_columns(_scholarship(PHD_CONFIGS), records, single_column_label=AWARD_COUNT_LABEL)
    payload = build_workbook(
        title="115學年度博士生獎學金 統計各學院學生獲獎人數",
        columns=columns,
        rows=tally_college_counts(records, columns),
        sheet_name=AWARD_REPORT_NAME,
    )
    ws = load_workbook(io.BytesIO(payload)).active
    assert ws.title == AWARD_REPORT_NAME
    assert [ws.cell(row=2, column=c).value for c in range(1, 5)] == ["編號", "學院", "國科會", "教育部"]
    by_label = {
        ws.cell(row=r, column=2).value: (ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value)
        for r in range(3, ws.max_row + 1)
    }
    assert by_label["資訊學院"] == (1, 1)
    assert by_label["電機學院"] == (0, 1)
    assert by_label[TOTAL_ROW_LABEL] == (1, 2)


def test_scholarship_columns_label_deactivated_sub_types_and_single_column():
    scholarship = _scholarship([("nstc", "x")], inactive=[("custom_grant", "自訂獎助")])
    columns = scholarship_stats_columns(
        scholarship, [_rec(1, "C", "custom_grant")], single_column_label=AWARD_COUNT_LABEL
    )
    assert [c.label for c in columns] == ["國科會", "自訂獎助"]
    assert scholarship_stats_columns(_scholarship([]), [_rec(1, "C")], single_column_label=AWARD_COUNT_LABEL) == [
        StatsColumn(SINGLE_COLUMN_KEY, AWARD_COUNT_LABEL)
    ]
