"""Admins and colleges see each application's 梯次 (多梯次申請 follow-up).

Label rule everywhere: a renewal is 續領, any other application 第N梯次 — 第一梯次
included. The admin list / detail payloads carry ``application_round`` and the
ranking, 彙整表 / 申請總表 and recipient exports carry a 申請梯次 column.
"""

import io
from types import SimpleNamespace

import pytest_asyncio
from openpyxl import load_workbook

from app.models.user import UserRole
from app.services.college_ranking_export_service import CollegeRankingExportService, ExportRow
from app.services.manual_distribution_export_service import build_recipient_row
from app.services.ranking_groups import application_round_label
from app.tests.ranking_scope_helpers import add_application, make_config, make_scholarship, make_user


def test_label_rule():
    assert application_round_label(SimpleNamespace(is_renewal=True, application_round=2)) == "續領"
    assert application_round_label(SimpleNamespace(is_renewal=False, application_round=1)) == "第一梯次"
    assert application_round_label(SimpleNamespace(is_renewal=False, application_round=2)) == "第二梯次"


def _app(**fields):
    return SimpleNamespace(
        sub_type_preferences=["nstc"],
        sub_scholarship_type="nstc",
        student_data={"std_cname": "某生"},
        submitted_form_data={},
        **fields,
    )


def test_ranking_export_has_a_round_column():
    rows = [
        ExportRow(rank_position=1, application=_app(is_renewal=True, application_round=1)),
        ExportRow(rank_position=2, application=_app(is_renewal=False, application_round=1)),
        ExportRow(rank_position=3, application=_app(is_renewal=False, application_round=2)),
    ]
    blob = CollegeRankingExportService().build_workbook(
        rows=rows, dynamic_fields=[], sub_type_labels={}, title="T", sheet_name="S"
    )
    ws = load_workbook(io.BytesIO(blob))["S"]
    headers = [cell.value for cell in ws[2]]
    column = headers.index("申請梯次") + 1
    assert [ws.cell(row=r, column=column).value for r in (3, 4, 5)] == ["續領", "第一梯次", "第二梯次"]


def test_ranking_export_pdf_renders_with_the_round_column():
    pdf = CollegeRankingExportService().build_pdf(
        rows=[ExportRow(rank_position=1, application=_app(is_renewal=False, application_round=2))],
        dynamic_fields=[],
        sub_type_labels={},
        title="T",
    )
    assert pdf.startswith(b"%PDF")


def test_recipient_export_row_carries_the_round():
    row = build_recipient_row(1, _app(is_renewal=False, application_round=2))
    assert row.application_round == "第二梯次"


class TestPayloads:
    @pytest_asyncio.fixture
    async def login(self):
        from app.core.security import get_current_user
        from app.main import app

        def _login(user) -> None:
            async def _override():
                return user

            app.dependency_overrides[get_current_user] = _override

        yield _login
        app.dependency_overrides.pop(get_current_user, None)

    @pytest_asyncio.fixture
    async def scholarship(self, db):
        return await make_scholarship(db, "visibility_scholarship")

    async def test_admin_list_and_detail_carry_round_and_renewal(self, client, db, login, scholarship):
        config = await make_config(db, scholarship.id, code="vis_cfg")
        second = await add_application(db, scholarship.id, "v2", application_round=2, configuration_id=config.id)
        renewal = await add_application(db, scholarship.id, "v0", is_renewal=True, configuration_id=config.id)
        login(await make_user(db, "vis_super", role=UserRole.super_admin))

        listing = await client.get(f"/api/v1/admin/scholarships/{scholarship.code}/applications")
        assert listing.status_code == 200, listing.text
        rows = {row["id"]: row for row in listing.json()["data"]}
        assert (rows[second.id]["application_round"], rows[second.id]["is_renewal"]) == (2, False)
        assert rows[renewal.id]["is_renewal"] is True

        detail = await client.get(f"/api/v1/applications/{second.id}")
        assert detail.status_code == 200, detail.text
        assert detail.json()["data"]["application_round"] == 2
