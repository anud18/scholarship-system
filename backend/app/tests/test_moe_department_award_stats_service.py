"""Unit tests for the 統計教育部博士生獎學金-各系所獲獎生人數 tally and layout."""

import io

from openpyxl import load_workbook

from app.services.moe_department_award_stats_service import (
    HEADERS,
    TOTAL_ROW_LABEL,
    UNKNOWN_DEPARTMENT_LABEL,
    UNMAPPED_FIELD_LABEL,
    build_workbook,
    department_code_from_data,
    department_record,
    tally_department_counts,
)
from app.services.moe_department_stats_layout import MOE_DEPARTMENT_FIELD_LAYOUT

NAMES = {"4460": "教育研究所", "4551": "資訊科學與工程研究所", "103": "電機資訊國際博士學位學程"}


def _record(user_id, depno, *sub_types, extra=None):
    return department_record(user_id, {"std_depno": depno, **(extra or {})}, sub_types, NAMES)


def _by_department(rows):
    return {row.department: row for row in rows}


def test_layout_matches_reference_form():
    labels = [label for _, rows in MOE_DEPARTMENT_FIELD_LAYOUT for label, _ in rows]
    assert len(labels) == 65
    assert len(set(labels)) == len(labels)
    assert [field for field, _ in MOE_DEPARTMENT_FIELD_LAYOUT] == [
        "01教育",
        "02藝術及人文領域",
        "04商業、管理及法律領域",
        "05自然科學、數學及統計領域",
        "06資訊通訊科技",
        "07工程、製造及營建領域",
        "09醫藥衛生及社會福利領域",
    ]


def test_counts_only_moe_sub_types_as_distinct_students():
    applicants = [
        _record(1, "4460", "moe_1w"),
        _record(1, "4460", "moe_2w"),  # same student twice ⇒ counted once
        _record(2, "4460", "nstc", "moe_2w"),
        _record(3, "4460", "nstc"),  # no 教育部 sub-type ⇒ not counted
    ]
    awardees = [_record(1, "4460", "moe_1w"), _record(2, "4460", "nstc")]

    row = _by_department(tally_department_counts(applicants, awardees))["教育研究所"]

    assert (row.field, row.applicants, row.awardees) == ("01教育", 2, 1)


def test_awardee_counts_as_applicant_even_without_the_moe_code_on_their_application():
    # An empty / placeholder sub-type list lets the distribution allocate any sub-type.
    applicants = [_record(1, "4460", "general")]
    awardees = [_record(1, "4460", "moe_1w")]

    row = _by_department(tally_department_counts(applicants, awardees))["教育研究所"]

    assert (row.applicants, row.awardees) == (1, 1)


def test_every_layout_row_is_emitted_and_unmapped_departments_follow():
    rows = tally_department_counts(
        [_record(1, "103", "moe_1w"), _record(2, "9999", "moe_1w"), _record(3, None, "moe_1w")], []
    )

    layout_count = sum(len(r) for _, r in MOE_DEPARTMENT_FIELD_LAYOUT)
    assert len(rows) == layout_count + 3
    tail = rows[layout_count:]
    assert {row.field for row in tail} == {UNMAPPED_FIELD_LABEL}
    assert sorted(row.department for row in tail) == sorted(
        ["電機資訊國際博士學位學程", f"{UNKNOWN_DEPARTMENT_LABEL}（9999）", UNKNOWN_DEPARTMENT_LABEL]
    )


def test_sis_name_variants_land_on_the_form_row():
    names = {"4307": "跨領域神經科學博士學位學程 (台灣聯合大學系統)", "4110": "環境與職業衛生研究所"}
    applicants = [
        department_record(1, {"std_depno": "4307"}, ["moe_1w"], names),
        department_record(2, {"std_depno": "4110"}, ["moe_2w"], names),
    ]

    by_department = _by_department(tally_department_counts(applicants, []))

    assert by_department["跨領域神經科學博士學位學程"].applicants == 1
    assert by_department["環境與職業衛生研究所"].applicants == 1


def test_department_name_prefers_departments_table_then_matching_term_snapshot():
    from_table = _record(1, "4460", "moe_1w", extra={"trm_depno": "4460", "trm_depname": "舊名"})
    from_term = department_record(2, {"trm_depno": "4551", "trm_depname": "資訊科學與工程研究所"}, ["moe_1w"], {})
    mismatched_term = department_record(
        3, {"std_depno": "7777", "trm_depno": "4551", "trm_depname": "X"}, ["moe_1w"], {}
    )

    assert from_table.department_name == "教育研究所"
    assert from_term.department_name == "資訊科學與工程研究所"
    assert mismatched_term.department_name is None
    assert mismatched_term.department_code == "7777"


def test_department_code_is_api_1_first_and_stripped():
    assert department_code_from_data({"std_depno": " 4460 ", "trm_depno": "4551"}) == "4460"
    assert department_code_from_data({"trm_depno": "4551"}) == "4551"  # term lookup only
    assert department_code_from_data({}) is None
    assert department_code_from_data(None) is None


def test_grouping_is_unchanged_when_std_and_trm_agree():
    agreeing = {"std_depno": "4460", "trm_depno": "4460", "trm_depname": "教育研究所"}
    std_only = {"std_depno": "4460"}
    term_only = {"trm_depno": "4460", "trm_depname": "教育研究所"}

    records = [department_record(i, data, ["moe_1w"], NAMES) for i, data in enumerate([agreeing, std_only, term_only])]
    row = _by_department(tally_department_counts(records, []))["教育研究所"]

    assert {r.department_code for r in records} == {"4460"}
    assert row.applicants == 3


def test_workbook_merges_field_cells_and_leaves_bonus_column_blank():
    rows = tally_department_counts([_record(1, "4460", "moe_1w")], [_record(1, "4460", "moe_1w")])

    ws = load_workbook(io.BytesIO(build_workbook(title="T", rows=rows))).active

    assert ws["A1"].value == "T"
    assert tuple(ws.cell(row=2, column=c).value for c in range(1, 6)) == HEADERS
    assert (ws["A3"].value, ws["B3"].value, ws["C3"].value, ws["D3"].value, ws["E3"].value) == (
        "01教育",
        "教育研究所",
        1,
        1,
        None,
    )
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:E1" in merged
    assert "A4:A6" in merged  # 02藝術及人文領域 spans three 系所, as on the form
    total_row = 3 + len(rows)
    assert ws.cell(row=total_row, column=1).value == TOTAL_ROW_LABEL
    assert (ws.cell(row=total_row, column=3).value, ws.cell(row=total_row, column=4).value) == (1, 1)
