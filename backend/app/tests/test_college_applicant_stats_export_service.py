"""Pure-logic tests for the 統計各學院學生申請人數 export service."""

import io

from openpyxl import load_workbook

from app.services.college_applicant_stats_export_service import (
    ALL_APPLICANTS_COLUMN,
    COLLEGE_STATS_LAYOUT,
    TOTAL_ROW_LABEL,
    UNKNOWN_COLLEGE_LABEL,
    UNSPECIFIED_SUB_TYPE_COLUMN,
    ApplicantRecord,
    StatsColumn,
    build_stats_columns,
    build_workbook,
    tally_college_counts,
)

PHD_CONFIGS = [("nstc", "國科會博士生獎學金"), ("moe_1w", "教育部 1 萬"), ("moe_2w", "教育部 2 萬")]


def _rec(user_id, academy, *sub_types):
    return ApplicantRecord(user_id=user_id, academy_code=academy, sub_type_codes=tuple(sub_types))


def _counts(rows):
    return {row.label: row.counts for row in rows}


def test_moe_sub_types_collapse_into_one_agency_column():
    columns = build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "nstc")])
    assert [c.label for c in columns] == ["國科會", "教育部"]


def test_ungrouped_sub_type_uses_its_configured_name():
    columns = build_stats_columns([("nstc", "x"), ("custom_new", "自訂獎助")], [])
    assert [c.label for c in columns] == ["國科會", "自訂獎助"]


def test_unconfigured_sub_type_on_an_application_still_gets_a_column():
    columns = build_stats_columns([("nstc", "x")], [_rec(1, "C", "legacy_type")])
    assert [c.key for c in columns] == ["國科會", "legacy_type"]


def test_scholarship_without_sub_types_gets_single_applicant_column():
    columns = build_stats_columns([], [_rec(1, "C", "general")])
    assert columns == [StatsColumn(ALL_APPLICANTS_COLUMN, ALL_APPLICANTS_COLUMN)]
    rows = tally_college_counts([_rec(1, "C", "general"), _rec(2, "C")], columns)
    assert _counts(rows)["資訊學院"] == (2,)


def test_unspecified_column_only_when_an_application_lacks_a_sub_type():
    assert UNSPECIFIED_SUB_TYPE_COLUMN not in [c.key for c in build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "nstc")])]
    columns = build_stats_columns(PHD_CONFIGS, [_rec(1, "C", "general")])
    assert columns[-1].key == UNSPECIFIED_SUB_TYPE_COLUMN
    assert _counts(tally_college_counts([_rec(1, "C", "general")], columns))["資訊學院"] == (0, 0, 1)


def test_tally_counts_each_student_once_per_column():
    records = [
        _rec(1, "C", "nstc", "moe_1w"),  # both agencies
        _rec(2, "C", "moe_1w", "moe_2w"),  # two MOE sub-types ⇒ one 教育部 count
        _rec(3, "E", "nstc"),
        _rec(3, "E", "nstc"),  # duplicate application, same student
    ]
    columns = build_stats_columns(PHD_CONFIGS, records)
    counts = _counts(tally_college_counts(records, columns))
    assert counts["資訊學院"] == (1, 2)
    assert counts["電機學院"] == (1, 0)
    assert counts[TOTAL_ROW_LABEL] == (2, 2)


def test_every_layout_college_is_listed_in_reference_order_even_at_zero():
    rows = tally_college_counts([], build_stats_columns(PHD_CONFIGS, []))
    assert [r.label for r in rows] == [label for label, _ in COLLEGE_STATS_LAYOUT] + [TOTAL_ROW_LABEL]
    assert all(r.counts == (0, 0) for r in rows)


def test_yangming_humanities_code_8_shares_the_humanities_row():
    records = [_rec(1, "A", "nstc"), _rec(2, "8", "nstc")]
    counts = _counts(tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records)))
    assert counts["人文藝術與社會學院"] == (2, 0)


def test_colleges_outside_layout_are_appended_before_total():
    records = [_rec(1, "Y", "nstc"), _rec(2, None, "nstc"), _rec(3, "~", "nstc")]
    rows = tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records))
    tail = [r.label for r in rows[len(COLLEGE_STATS_LAYOUT) :]]
    assert tail == ["電資學院", f"{UNKNOWN_COLLEGE_LABEL}（~）", UNKNOWN_COLLEGE_LABEL, TOTAL_ROW_LABEL]
    assert rows[-1].counts == (3, 0)


def test_workbook_matches_reference_layout():
    records = [_rec(1, "I", "nstc"), _rec(2, "I", "moe_1w")]
    columns = build_stats_columns(PHD_CONFIGS, records)
    payload = build_workbook(
        title="115學年度博士生獎學金 統計各學院學生申請人數",
        columns=columns,
        rows=tally_college_counts(records, columns),
    )

    ws = load_workbook(io.BytesIO(payload)).active
    assert ws["A1"].value == "115學年度博士生獎學金 統計各學院學生申請人數"
    assert "A1:D1" in {str(r) for r in ws.merged_cells.ranges}
    assert [ws.cell(row=2, column=c).value for c in range(1, 5)] == ["編號", "學院", "國科會", "教育部"]
    assert [ws.cell(row=4, column=c).value for c in range(1, 5)] == [2, "工學院", 1, 1]
    last = 2 + len(COLLEGE_STATS_LAYOUT) + 1
    assert [ws.cell(row=last, column=c).value for c in range(1, 5)] == [None, TOTAL_ROW_LABEL, 1, 1]
    assert ws["B3"].font.name == "微軟正黑體"


def test_workbook_neutralises_formula_like_sub_type_names():
    columns = build_stats_columns([("custom", "=HYPERLINK(1)")], [])
    payload = build_workbook(title="t", columns=columns, rows=tally_college_counts([], columns))
    ws = load_workbook(io.BytesIO(payload)).active
    assert not str(ws["C2"].value).startswith("=")


def test_total_row_is_the_column_sum_even_when_a_student_spans_two_colleges():
    records = [_rec(1, "C", "nstc"), _rec(1, "E", "nstc")]  # college changed between applications
    counts = _counts(tally_college_counts(records, build_stats_columns(PHD_CONFIGS, records)))
    assert counts["資訊學院"] == (1, 0) and counts["電機學院"] == (1, 0)
    assert counts[TOTAL_ROW_LABEL] == (2, 0)


def test_deactivated_sub_type_column_keeps_its_configured_name():
    columns = build_stats_columns(
        [("nstc", "x")],
        [_rec(1, "C", "custom_grant")],
        known_names={"nstc": "x", "custom_grant": "自訂獎助"},
    )
    assert [c.label for c in columns] == ["國科會", "自訂獎助"]
