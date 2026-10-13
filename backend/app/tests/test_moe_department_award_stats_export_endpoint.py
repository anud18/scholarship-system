"""Endpoint tests for 統計教育部博士生獎學金-各系所獲獎生人數.

GET /api/v1/manual-distribution/distribution-summary/moe-department-award-stats/export

Counts are asserted by parsing the workbook, never by scanning response bytes.
"""

import io
from urllib.parse import unquote

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_admin_user, get_current_user
from app.main import app
from app.models.application import Application
from app.models.college_review import CollegeRanking, CollegeRankingItem
from app.models.enums import ApplicationStatus, ReviewStage, Semester, SubTypeSelectionMode
from app.models.scholarship import ScholarshipSubTypeConfig, ScholarshipType
from app.models.student import Department
from app.models.user import User, UserRole, UserType
from app.services.moe_department_award_stats_service import TOTAL_ROW_LABEL, UNMAPPED_FIELD_LABEL
from app.services.stats_applicant_source import load_applicant_rows

ACADEMIC_YEAR = 114
URL = "/api/v1/manual-distribution/distribution-summary/moe-department-award-stats/export"
PARAMS = {"academic_year": ACADEMIC_YEAR, "semester": "yearly"}


async def _user(db: AsyncSession, nycu_id: str, role: UserRole) -> User:
    user = User(
        nycu_id=nycu_id,
        name=nycu_id,
        email=f"{nycu_id}@university.edu",
        user_type=UserType.student if role == UserRole.student else UserType.employee,
        role=role,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


@pytest_asyncio.fixture
async def admin_client(client: AsyncClient, db: AsyncSession):
    admin = await _user(db, "moe_dept_admin", UserRole.admin)
    app.dependency_overrides[get_current_admin_user] = lambda: admin
    app.dependency_overrides[get_current_user] = lambda: admin
    try:
        yield client
    finally:
        app.dependency_overrides.pop(get_current_admin_user, None)
        app.dependency_overrides.pop(get_current_user, None)


async def _scholarship(db: AsyncSession, code: str, sub_types) -> ScholarshipType:
    st = ScholarshipType(code=code, name="博士生獎學金")
    db.add(st)
    await db.commit()
    await db.refresh(st)
    for order, sub_type in enumerate(sub_types, 1):
        db.add(
            ScholarshipSubTypeConfig(
                scholarship_type_id=st.id, sub_type_code=sub_type, name=sub_type, display_order=order, is_active=True
            )
        )
    await db.commit()
    return st


@pytest_asyncio.fixture
async def phd(db: AsyncSession) -> ScholarshipType:
    for code, name in [
        ("4460", "教育研究所"),
        ("4551", "資訊科學與工程研究所"),
        ("4354", "跨領域神經科學博士學位學程(台灣聯合大學系統)"),
        ("103", "電機資訊國際博士學位學程"),
    ]:
        db.add(Department(code=code, name=name))
    await db.commit()
    return await _scholarship(db, "moe_dept_phd", ["nstc", "moe_1w", "moe_2w"])


async def _application(
    db: AsyncSession,
    st: ScholarshipType,
    suffix: str,
    depno: str,
    sub_types,
    status=ApplicationStatus.submitted,
    *,
    is_renewal: bool = False,
) -> Application:
    user = await _user(db, f"moe_dept_stu_{suffix}", UserRole.student)
    application = Application(
        app_id=f"APP-{ACADEMIC_YEAR}-0-7{suffix}",
        user_id=user.id,
        scholarship_type_id=st.id,
        scholarship_subtype_list=list(sub_types),
        sub_type_selection_mode=SubTypeSelectionMode.multiple,
        sub_scholarship_type=sub_types[0],
        academic_year=ACADEMIC_YEAR,
        semester=None,
        status=status,
        review_stage=ReviewStage.quota_distributed,
        agree_terms=True,
        is_renewal=is_renewal,
        student_data={"std_stdcode": f"31{suffix}", "std_academyno": "C", "std_depno": depno},
    )
    db.add(application)
    await db.commit()
    await db.refresh(application)
    return application


async def _seed_distribution(db: AsyncSession, st: ScholarshipType, rows, *, executed: bool = True) -> None:
    """``rows``: (suffix, depno, applied sub-types, allocated sub-type or None, status)."""
    ranking = CollegeRanking(
        scholarship_type_id=st.id,
        sub_type_code="default",
        academic_year=ACADEMIC_YEAR,
        semester=None,
        is_finalized=True,
        distribution_executed=executed,
        ranking_status="finalized",
    )
    db.add(ranking)
    await db.commit()
    await db.refresh(ranking)
    for rank, (suffix, depno, sub_types, allocated, status) in enumerate(rows, 1):
        application = await _application(db, st, suffix, depno, sub_types, status)
        db.add(
            CollegeRankingItem(
                ranking_id=ranking.id,
                application_id=application.id,
                rank_position=rank,
                is_allocated=allocated is not None,
                allocated_sub_type=allocated,
                status="allocated" if allocated else "ranked",
            )
        )
    await db.commit()


SUBMITTED = ApplicationStatus.submitted
ROWS = [
    ("01", "4460", ["moe_1w"], "moe_1w", SUBMITTED),
    ("02", "4551", ["nstc", "moe_2w"], "nstc", SUBMITTED),  # 教育部 applicant, 國科會 award
    ("03", "4551", ["moe_1w"], None, SUBMITTED),  # ranked, not allocated
    ("04", "4354", ["moe_2w"], "moe_2w", SUBMITTED),  # SIS name variant of a form row
    ("05", "103", ["moe_1w"], "moe_1w", SUBMITTED),  # 系所 missing from the form
    ("06", "4460", ["nstc"], "nstc", SUBMITTED),  # 國科會 only ⇒ not counted
    ("07", "4460", ["moe_1w"], None, ApplicationStatus.draft),  # draft ⇒ not an applicant
    ("08", "4460", ["general"], "moe_2w", SUBMITTED),  # allocated 教育部 without listing it ⇒ applicant too
]


def _rows(content: bytes):
    ws = load_workbook(io.BytesIO(content)).active
    body = {
        ws.cell(row=r, column=2).value: (ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value)
        for r in range(3, ws.max_row)
    }
    total = (ws.cell(row=ws.max_row, column=3).value, ws.cell(row=ws.max_row, column=4).value)
    return ws, body, total


@pytest.mark.asyncio
async def test_counts_moe_applicants_and_awardees_per_department(admin_client, db, phd):
    await _seed_distribution(db, phd, ROWS)
    # 續領 never enters the 甄選 ⇒ not a 申請人數.
    await _application(db, phd, "09", "4460", ["moe_1w"], ApplicationStatus.approved, is_renewal=True)

    resp = await admin_client.get(URL, params={"scholarship_type_id": phd.id, **PARAMS})

    assert resp.status_code == 200, resp.text
    filename = unquote(resp.headers["content-disposition"])
    assert "114學年度統計教育部博士生獎學金-各系所獲獎生人數.xlsx" in filename
    ws, body, total = _rows(resp.content)
    assert ws["A1"].value == "114學年度統計教育部博士生獎學金-各系所獲獎生人數"
    assert body["教育研究所"] == (2, 2)
    assert body["資訊科學與工程研究所"] == (2, 0)
    assert body["跨領域神經科學博士學位學程"] == (1, 1)
    assert body["電機資訊國際博士學位學程"] == (1, 1)
    assert body["應用藝術研究所"] == (0, 0)
    assert ws.cell(row=ws.max_row - 1, column=1).value == UNMAPPED_FIELD_LABEL
    assert ws.cell(row=ws.max_row, column=1).value == TOTAL_ROW_LABEL
    assert total == (6, 4)


@pytest.mark.asyncio
async def test_400_when_scholarship_has_no_moe_sub_type(admin_client, db):
    st = await _scholarship(db, "moe_dept_nstc_only", ["nstc"])
    await _seed_distribution(db, st, [("11", "4460", ["nstc"], "nstc", SUBMITTED)])

    resp = await admin_client.get(URL, params={"scholarship_type_id": st.id, **PARAMS})

    assert resp.status_code == 400
    assert "教育部" in resp.json()["message"]


@pytest.mark.asyncio
async def test_404_before_distribution_runs(admin_client, db, phd):
    await _seed_distribution(db, phd, ROWS[:1], executed=False)

    resp = await admin_client.get(URL, params={"scholarship_type_id": phd.id, **PARAMS})

    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_404_for_unknown_scholarship(admin_client):
    resp = await admin_client.get(URL, params={"scholarship_type_id": 999999, **PARAMS})

    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_rejects_unknown_semester(admin_client, phd):
    resp = await admin_client.get(
        URL, params={"scholarship_type_id": phd.id, "academic_year": 114, "semester": "summer"}
    )

    assert resp.status_code == 422


@pytest.mark.asyncio
@pytest.mark.parametrize("semester", ["yearly", "annual", "Yearly"])
async def test_applicant_loader_treats_every_yearly_spelling_as_null_semester(db, phd, semester):
    await _application(db, phd, "21", "4460", ["moe_1w"])  # semester NULL
    stored_yearly = await _application(db, phd, "22", "4460", ["moe_2w"])
    stored_yearly.semester = Semester.yearly
    await db.commit()
    first = await _application(db, phd, "23", "4460", ["nstc"])
    first.semester = Semester.first
    await db.commit()

    rows = await load_applicant_rows(db, scholarship_type_id=phd.id, academic_year=ACADEMIC_YEAR, semester=semester)

    assert sorted(row.sub_type_codes for row in rows) == [("moe_1w",), ("moe_2w",)]
