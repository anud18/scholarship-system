"""Build the two AY115 sample import workbooks that pair with the dev seed.

    docs/samples/ay115/115_續領生匯入.xlsx   → admin 匯入續領生 (phd / 115)
    docs/samples/ay115/115_批次匯入.xlsx     → admin 批次匯入   (phd / 115)

The rows mirror `backend/app/db/seed_ay115_demo.py` (續領 sheet = the 114
recipients) and `mock-student-api/main.py::_AY115_DEMO` (批次匯入 sheet =
students with SIS rows but no account yet). Column headers are the exact
strings the two parsers read (`renewal_import_service.parse_renewal_excel`
and `batch_import_template_service.build_batch_import_template`).

Usage (host, needs openpyxl):
    python backend/scripts/generate_ay115_import_samples.py [output_dir]
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[2] / "docs" / "samples" / "ay115"

ADVISORS = {
    "C": ("李資訊教授", "cs_professor@nycu.edu.tw", "cs_professor"),
    "E": ("李教授", "professor@nycu.edu.tw", "professor"),
}

# 續領生匯入: 編號, 學院, 系所, 學生姓名, 學號, 學生年級, 學生是否申請續領, 續領審核結果,
#             獎學金類別, 郵局帳號, 指導教授本校人事編號, 指導教授姓名
RENEWAL_HEADERS = [
    "編號",
    "學院",
    "系所",
    "學生姓名",
    "學號",
    "學生年級",
    "學生是否申請續領",
    "續領審核結果",
    "獎學金類別",
    "郵局帳號",
    "指導教授本校人事編號",
    "指導教授姓名",
]
# (學院代碼, 學院, 系所, 姓名, 學號, 年級, 是否續領, 審核結果, 類別)
RENEWAL_ROWS = [
    ("C", "資訊學院", "資訊工程學系", "續領-林承翰", "313551201", "博三", "是", "通過", "國科會"),
    ("C", "資訊學院", "資訊科學與工程研究所", "續領-張雅婷", "313551202", "博三", "是", "通過", "教育部"),
    ("E", "電機學院", "電機工程學系", "續領-黃冠宇", "312551203", "博四", "是", "通過", "國科會"),
    ("E", "電機學院", "電機學院博士班", "續領-吳佩珊", "313551204", "博三", "是", "通過", "教育部"),
    ("C", "資訊學院", "資訊工程學系", "續領-鄭宇軒", "311551205", "博五", "否", "領獎期滿，無續領", "國科會"),
    ("C", "資訊學院", "資訊科學與工程研究所", "114新申請-何冠廷", "314551206", "博二", "是", "通過", "國科會"),
    ("E", "電機學院", "電機工程學系", "114新申請-謝宜庭", "314551207", "博二", "是", "通過", "教育部"),
]

# 批次匯入: 學號, 學生姓名, 郵局帳號, 指導教授姓名, 指導教授Email, 指導教授本校人事編號,
#           國科會, 教育部, 碩士畢業學校/學院/系所
BATCH_HEADERS = [
    "學號",
    "學生姓名",
    "郵局帳號",
    "指導教授姓名",
    "指導教授Email",
    "指導教授本校人事編號",
    "國科會",
    "教育部",
    "碩士畢業學校/學院/系所",
]
# (學號, 姓名, 學院代碼, 國科會, 教育部, 碩士畢業)
BATCH_ROWS = [
    ("314551301", "匯入-陳柏宇", "C", 1, 1, "國立陽明交通大學資訊學院資訊工程學系"),
    ("314551302", "匯入-李欣怡", "C", 0, 1, "國立臺灣大學電機資訊學院資訊工程學系"),
    ("314551303", "匯入-王志豪", "E", 1, 0, "國立清華大學電機資訊學院電機工程學系"),
    ("313551304", "匯入-劉育綺", "E", 1, 1, "國立成功大學電機資訊學院電機工程學系"),
]

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def _postal_account(student_id: str) -> str:
    return f"0021{student_id}"


def _renewal_sheet_rows():
    for idx, (college_code, college, dept, name, sid, grade, applied, result, category) in enumerate(
        RENEWAL_ROWS, start=1
    ):
        advisor_name, _email, advisor_id = ADVISORS[college_code]
        renews = applied == "是"
        yield [
            idx,
            college,
            dept,
            name,
            sid,
            grade,
            applied,
            result,
            category,
            _postal_account(sid) if renews else "",
            advisor_id if renews else "",
            advisor_name if renews else "",
        ]


def _batch_sheet_rows():
    for sid, name, college_code, nstc, moe, master_school in BATCH_ROWS:
        advisor_name, advisor_email, advisor_id = ADVISORS[college_code]
        yield [sid, name, _postal_account(sid), advisor_name, advisor_email, advisor_id, nstc, moe, master_school]


def _write_workbook(path: Path, sheet_name: str, headers, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(list(row))
    for col_idx, header in enumerate(headers, start=1):
        values = [str(header)] + [
            str(sheet.cell(row=r, column=col_idx).value or "") for r in range(2, sheet.max_row + 1)
        ]
        widest = max(len(v) + sum(1 for ch in v if "一" <= ch <= "鿿") for v in values)
        sheet.column_dimensions[get_column_letter(col_idx)].width = widest + 2
    # Student ids / postal accounts must stay text so Excel never strips leading zeros.
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.isdigit():
                cell.number_format = "@"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main(output_dir: Path) -> None:
    renewal_path = output_dir / "115_續領生匯入.xlsx"
    batch_path = output_dir / "115_批次匯入.xlsx"
    _write_workbook(renewal_path, "續領生名單", RENEWAL_HEADERS, _renewal_sheet_rows())
    _write_workbook(batch_path, "批次匯入", BATCH_HEADERS, _batch_sheet_rows())
    print(f"wrote {renewal_path}")
    print(f"wrote {batch_path}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT_DIR)
